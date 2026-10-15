"""
Verify the Gotenberg PDF service (V1.6 Phase 2) — no database involved.

Usage (from backend/):
    python scripts/verify_gotenberg.py            # GET /health only
    python scripts/verify_gotenberg.py --live     # ALSO convert a tiny DOCX to PDF
    python scripts/verify_gotenberg.py --url http://localhost:3001

Where the URL comes from (first match wins):
  1. --url
  2. the GOTENBERG_URL environment variable
  3. GOTENBERG_URL in the repository-root .env (only that key is read)
  4. the default http://localhost:3000

What it checks (in order):
  1. GET {url}/health answers 200
  2. --live only: POST a tiny in-memory DOCX to {url}/forms/libreoffice/convert
     (multipart field "files", the same request the backend sends) and check
     that the answer is a PDF (starts with %PDF); prints size and latency.

Start Gotenberg with:  docker compose up -d gotenberg   (repo-root compose file)
Exit code 0 = everything it checked passed; 1 = something failed.
"""

import argparse
import io
import os
import sys
import time
from pathlib import Path

import httpx

BACKEND_DIR = Path(__file__).resolve().parents[1]
REPO_ENV = BACKEND_DIR.parent / ".env"
DEFAULT_URL = "http://localhost:3000"
DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
START_HINT = "Gotenberg is not running. Start it with: docker compose up -d gotenberg"

PASS = "  [OK]  "
FAIL = "  [FAIL]"
INFO = "  [INFO]"

problems = 0


def report(ok: bool, message: str) -> None:
    global problems
    print(f"{PASS if ok else FAIL} {message}")
    if not ok:
        problems += 1


def resolve_url(cli_url: str | None) -> tuple[str, str]:
    """(url, where it came from). Reads only GOTENBERG_URL from .env."""
    if cli_url:
        return cli_url.strip().rstrip("/"), "--url"
    from_env = (os.environ.get("GOTENBERG_URL") or "").strip()
    if from_env:
        return from_env.rstrip("/"), "environment"
    if REPO_ENV.is_file():
        try:
            from dotenv import dotenv_values

            from_file = (dotenv_values(REPO_ENV).get("GOTENBERG_URL") or "").strip()
        except Exception:  # unreadable .env: fall back to the default
            from_file = ""
        if from_file:
            return from_file.rstrip("/"), "repo .env"
    return DEFAULT_URL, "default"


def check_health(url: str) -> bool:
    try:
        response = httpx.get(f"{url}/health", timeout=5)
    except httpx.ConnectError:
        report(False, f"health -> cannot connect to {url}. {START_HINT}")
        return False
    except httpx.HTTPError as exc:
        report(False, f"health -> {type(exc).__name__} from {url}/health")
        return False
    ok = response.status_code == 200
    detail = ""
    try:
        detail = f" status={response.json().get('status')}"
    except ValueError:
        pass
    report(ok, f"health -> HTTP {response.status_code}{detail}")
    return ok


def _tiny_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("TemplateOS Gotenberg check", level=1)
    document.add_paragraph("If you can read this in a PDF, conversion works.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def live_convert(url: str, timeout: float) -> bool:
    docx_bytes = _tiny_docx()
    start = time.perf_counter()
    try:
        response = httpx.post(
            f"{url}/forms/libreoffice/convert",
            files={"files": ("verify-gotenberg.docx", docx_bytes, DOCX_MEDIA_TYPE)},
            timeout=httpx.Timeout(timeout, connect=5.0),
        )
    except httpx.ConnectError:
        report(False, f"convert -> cannot connect to {url}. {START_HINT}")
        return False
    except httpx.TimeoutException:
        report(False, f"convert -> timed out after {timeout:.0f}s")
        return False
    except httpx.HTTPError as exc:
        report(False, f"convert -> {type(exc).__name__}")
        return False
    elapsed = time.perf_counter() - start
    is_pdf = response.content.startswith(b"%PDF")
    ok = response.status_code == 200 and is_pdf
    trace = response.headers.get("Gotenberg-Trace", "-")
    message = (
        f"convert -> HTTP {response.status_code}, {len(response.content)} bytes, "
        f"{elapsed:.2f}s, {'PDF' if is_pdf else 'NOT a PDF'} (DOCX in: {len(docx_bytes)} bytes, "
        f"trace {trace})"
    )
    report(ok, message)
    return ok


def main() -> int:
    parser = argparse.ArgumentParser(description="Verify the Gotenberg PDF service.")
    parser.add_argument("--live", action="store_true", help="also convert a tiny DOCX to PDF")
    parser.add_argument("--url", help="Gotenberg base URL (overrides GOTENBERG_URL)")
    parser.add_argument(
        "--timeout", type=float, default=60.0, help="conversion timeout in seconds (default 60)"
    )
    args = parser.parse_args()

    url, source = resolve_url(args.url)
    print("Gotenberg PDF service check")
    print(f"{INFO} url: {url} (from {source})")

    if check_health(url) and args.live:
        live_convert(url, args.timeout)
    elif not args.live:
        print(f"{INFO} conversion not tested (add --live)")

    if problems:
        print(f"\n{problems} check(s) failed.")
        return 1
    print("\nAll checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
