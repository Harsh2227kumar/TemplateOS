"""
Smart template pipeline tests: cleaning suggestions -> anchored cleaning ->
placeholder detection -> field inference -> field editor robustness.

Every fixture is a realistic DOCX built in memory with python-docx (plus a
hand-written Word text box with its VML fallback twin). No network, no AWS:
the database is in-memory SQLite and storage is the local test folder.
"""

import io
import json
import os
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docxtpl import DocxTemplate
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault(
    "JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests"
)

from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.template import Template
from app.models.user import User
from app.services import docx_parser
from app.services.clean_suggester import suggest_cleaning
from app.services.docx_cleaner import apply_replacements_with_results, check_renderable
from app.services.field_inference import (
    infer_field_meta,
    infer_field_type,
    key_to_label,
    label_to_key,
    suggest_key,
)

test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestSession = sessionmaker(bind=test_engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    old = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


# ── fixtures: realistic university documents ────────────────────────

TEXTBOX_RUN = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
     xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing>
        <wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1"
                   behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>
          <wp:extent cx="1828800" cy="457200"/>
          <wp:wrapNone/>
          <wp:docPr id="10" name="Signatory"/>
          <a:graphic>
            <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
              <wps:wsp>
                <wps:cNvSpPr txBox="1"/>
                <wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="1828800" cy="457200"/></a:xfrm>
                  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>
                <wps:txbx>
                  <w:txbxContent>
                    <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Dr. Asha Verma</w:t></w:r></w:p>
                    <w:p><w:r><w:t>Head of Department, CSE</w:t></w:r></w:p>
                  </w:txbxContent>
                </wps:txbx>
                <wps:bodyPr/>
              </wps:wsp>
            </a:graphicData>
          </a:graphic>
        </wp:anchor>
      </w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict>
        <v:shape style="width:144pt;height:36pt">
          <v:textbox>
            <w:txbxContent>
              <w:p><w:r><w:rPr><w:b/></w:rPr><w:t>Dr. Asha Verma</w:t></w:r></w:p>
              <w:p><w:r><w:t>Head of Department, CSE</w:t></w:r></w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


def _bytes(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _merge_vertically(cell_top, cell_bottom) -> None:
    cell_top.merge(cell_bottom)


def build_notice_docx() -> bytes:
    """A department notice/form mixing every 'empty value' style."""
    document = Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "Ref. No: ________"
    section.footer.paragraphs[0].text = "Contact: hod.cse@univ.edu"

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("NOTICE").bold = True

    document.add_paragraph("Date:")
    document.add_paragraph("Venue .........")
    document.add_paragraph("Name:\t\t\t")
    document.add_paragraph("Department :")
    document.add_paragraph("Signature ____")
    document.add_paragraph("Student: [Name]  Roll No: <Roll Number>")
    document.add_paragraph("Branch {Department} for «Course» (Signature)")

    # Filled label/value pair split across runs: bold label, plain value.
    filled = document.add_paragraph()
    filled.add_run("Subject: ").bold = True
    filled.add_run("Annual ")
    filled.add_run("Sports Meet")

    document.add_paragraph(
        "All students are informed that the seminar will be held on 12/03/2024 "
        "at 10:30 AM. Queries: events@univ.edu or +91 98765 43210."
    )
    document.add_paragraph("Existing {{ already_there }} stays untouched.")
    return _bytes(document)


def build_table_docx() -> bytes:
    """Row-wise label/value form + column-header table with a merged cell."""
    document = Document()
    document.add_paragraph("Application Form").runs[0].bold = True

    form = document.add_table(rows=3, cols=2)
    form.cell(0, 0).text = "Name of the Student"
    form.cell(1, 0).text = "Date of Birth"
    form.cell(1, 1).text = "14/02/2003"
    form.cell(2, 0).text = "Email"

    document.add_paragraph("Attendance")
    grid = document.add_table(rows=3, cols=3)
    grid.cell(0, 0).text = "Agenda"
    grid.cell(0, 1).text = "Action Items"
    grid.cell(0, 2).text = "Target Date"
    # Horizontally merged cell in the last data row.
    grid.cell(2, 0).merge(grid.cell(2, 1))
    return _bytes(document)


def build_mom_docx() -> bytes:
    """Minutes of Meeting shaped like the real SIT Nagpur format."""
    document = Document()
    first = document.add_paragraph()
    first.add_run("Ref. No: SITN/2601/Nomenclature/Sr.No").bold = True
    first.add_run("\t\t\t")
    first.add_run("Date: _____________").bold = True
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("Minutes of Meeting").bold = True
    document.add_paragraph(
        "A meeting of the ------------- was held on --------- at the ---------, "
        "at ------ AM/PM. The meeting was held by Prof. -------------, "
        "Symbiosis Institute of Technology, Nagpur."
    )
    agenda = document.add_paragraph()
    agenda.add_run("Agenda Items").bold = True
    document.add_paragraph("ABC", style="List Number")
    document.add_paragraph("ABC", style="List Number")
    discussion = document.add_paragraph()
    discussion.add_run("Discussion").bold = True
    document.add_paragraph("The following points were discussed")
    for number in ("1.", "2.", "3. "):
        document.add_paragraph().add_run(number).bold = True
    document.add_paragraph().add_run("Prepared by:").bold = True
    signatory = document.add_paragraph()
    signatory._p.append(parse_xml(TEXTBOX_RUN))
    return _bytes(document)


def _accept_all(docx_bytes: bytes, min_confidence: float = 0.0):
    report = suggest_cleaning(docx_bytes)
    return report, [
        {
            "sample_text": s.sample_text,
            "placeholder_key": s.key,
            "anchor": {
                "paragraph_id": s.pid,
                "start": s.start,
                "end": s.end,
                "action": s.action,
                "paragraph_ids": s.paragraph_ids,
            },
        }
        for s in report.suggestions
        if s.confidence >= min_confidence
    ]


def _render(docx_bytes: bytes, context: dict) -> Document:
    template = DocxTemplate(io.BytesIO(docx_bytes))
    template.render(context)
    buffer = io.BytesIO()
    template.save(buffer)
    return Document(io.BytesIO(buffer.getvalue()))


def _all_text(document) -> str:
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    for section in document.sections:
        parts += [p.text for p in section.header.paragraphs]
        parts += [p.text for p in section.footer.paragraphs]
    return "\n".join(parts)


def _by_key(report):
    out = {}
    for suggestion in report.suggestions:
        out.setdefault(suggestion.key, []).append(suggestion)
    return out


# ── 1. field inference ───────────────────────────────────────────────


@pytest.mark.parametrize(
    "key,expected",
    [
        ("meeting_date", "date"),
        ("dob", "date"),
        ("submission_deadline", "date"),
        ("attendee_count", "number"),
        ("total_amount", "number"),
        ("marks", "number"),
        ("sr_no", "number"),
        ("ref_no", "text"),
        ("roll_no", "text"),
        ("attendees", "list"),
        ("agenda_items", "list"),
        ("discussion_points", "list"),
        ("members", "list"),
        ("agenda", "list"),
        ("description", "textarea"),
        ("meeting_minutes", "textarea"),
        ("remarks", "textarea"),
        ("event_summary", "textarea"),
        ("hod_signature", "signature"),
        ("sign", "signature"),
        ("venue", "text"),
        ("employee_name", "text"),
    ],
)
def test_infer_field_type(key, expected):
    assert infer_field_type(key) == expected


def test_infer_type_uses_sample_value():
    assert infer_field_type("ref_no", value="SITN/2601/Nomenclature") == "text"
    assert infer_field_type("held", value="12/03/2024") == "date"
    assert infer_field_type("agenda_1", allow_list=False) == "text"


def test_field_meta_label_validation_required_and_section():
    meta = infer_field_meta("contact_email", value="a@b.co", section="Details")
    assert meta.field_label == "Contact Email"
    assert meta.validation_rule == "email"
    assert meta.example_value == "a@b.co"
    assert meta.section == "Details"
    assert infer_field_meta("remarks").is_required is False
    assert infer_field_meta("prepared_by", section="Discussion").section == "Signatures"
    assert infer_field_meta("mobile_no").validation_rule == "min:10|max:20"
    # Lists are not rewritable by POST /ai/improve-text (text/textarea only).
    assert infer_field_meta("agenda_items").ai_enabled is False
    assert infer_field_meta("meeting_summary").ai_enabled is True
    assert infer_field_meta("meeting_date").ai_enabled is False


def test_keys_and_labels():
    assert label_to_key("Title of the Meeting") == "meeting_title"
    assert label_to_key("Ref. No") == "ref_no"
    assert label_to_key("E-mail ID") == "email_id"
    assert label_to_key("Date of Birth") == "date_of_birth"
    assert label_to_key("Name & Sign of Faculty") == "name_sign_of_faculty"
    assert key_to_label("hod_name") == "HOD Name"
    assert key_to_label("date_of_birth") == "Date of Birth"
    assert suggest_key("Meeting Title") == "meeting_title"
    assert suggest_key("meetingDate") == "meeting_date"
    assert suggest_key("2nd-date") == "field_2nd_date"


# ── 2. cleaning suggestions ──────────────────────────────────────────


def test_suggests_empty_values_leaders_and_brackets():
    report = suggest_cleaning(build_notice_docx())
    keys = _by_key(report)

    # Empty values in every spelling.
    for key in ("date", "name", "department"):
        assert key in keys, f"missing {key}: {sorted(keys)}"
    date = keys["date"][0]
    assert date.kind == "empty_value" and date.action == "insert"
    assert date.field_type == "date"
    assert date.confidence >= 0.8

    # Underscore / dot leaders keep the label as key.
    assert keys["venue"][0].kind == "leader"
    assert keys["venue"][0].sample_text == "........."
    assert keys["signature"][0].field_type == "signature"

    # Bracket hints of every kind.
    brackets = {s.sample_text for s in report.suggestions if s.kind == "bracket"}
    assert {"[Name]", "<Roll Number>", "{Department}", "«Course»", "(Signature)"} <= brackets

    # Repeated labels collapse onto one key: [Name] reuses "name",
    # {Department} reuses "department", (Signature) reuses "signature".
    assert len(keys["name"]) == 2
    assert len(keys["department"]) == 2
    assert all(s.shared_with == 2 for s in keys["name"])


def test_suggests_filled_pairs_patterns_headers_and_footers():
    report = suggest_cleaning(build_notice_docx())
    keys = _by_key(report)

    subject = keys["subject"][0]
    assert subject.kind == "filled_value"
    assert subject.sample_text == "Annual Sports Meet"
    assert subject.example_value == "Annual Sports Meet"

    # Header leader + footer email (label/value pair).
    ref = keys["ref_no"][0]
    assert ref.location == "header" and ref.section == "Header"
    contact = keys["contact"][0]
    assert contact.location == "footer"
    assert contact.validation_rule == "email"

    # Lower-confidence values in running text.
    detected = {s.key: s for s in report.suggestions if s.kind == "detected_value"}
    assert detected["seminar_date"].sample_text == "12/03/2024"
    assert detected["seminar_time"].sample_text == "10:30 AM"
    assert detected["email"].sample_text == "events@univ.edu"
    assert detected["phone"].sample_text == "+91 98765 43210"
    assert all(s.confidence < 0.8 for s in detected.values())


def test_never_suggests_inside_existing_placeholders():
    report = suggest_cleaning(build_notice_docx())
    assert report.existing_keys == ["already_there"]
    for suggestion in report.suggestions:
        assert "{{" not in suggestion.sample_text
        assert "already_there" not in suggestion.key


def test_suggests_table_cells_row_and_column_wise():
    report = suggest_cleaning(build_table_docx())
    keys = _by_key(report)

    assert keys["student_name"][0].kind == "table_cell"
    assert keys["student_name"][0].action == "insert"
    assert keys["date_of_birth"][0].kind == "table_value"
    assert keys["date_of_birth"][0].field_type == "date"
    assert keys["email"][0].validation_rule == "email"

    # Column-header table: one key per column per data row; the merged
    # cell yields one suggestion, not two.
    for key in ("agenda_1", "action_items_1", "target_date_1", "agenda_2", "target_date_2"):
        assert key in keys, sorted(keys)
    assert "action_items_2" not in keys
    assert keys["target_date_1"][0].field_type == "date"
    assert keys["agenda_1"][0].field_type == "text"  # not a list inside a cell


def test_suggests_lists_and_signatories_for_mom():
    report = suggest_cleaning(build_mom_docx())
    keys = _by_key(report)

    assert keys["ref_no"][0].sample_text == "SITN/2601/Nomenclature/Sr.No"
    assert keys["date"][0].kind == "leader"
    for key in ("committee_name", "meeting_date", "meeting_venue", "meeting_time", "chairperson_name"):
        assert keys[key][0].kind == "inline_blank", key
    assert keys["meeting_date"][0].field_type == "date"

    agenda = keys["agenda_items"][0]
    assert agenda.kind == "list" and agenda.field_type == "list"
    assert len(agenda.paragraph_ids) == 2
    assert agenda.section == "Agenda Items"
    points = keys["discussion_points"][0]
    assert points.kind == "list" and len(points.paragraph_ids) == 3

    assert keys["prepared_by"][0].action == "insert"
    assert keys["prepared_by"][0].section == "Signatures"

    hod = keys["head_of_department_name"][0]
    assert hod.location == "textbox"
    assert hod.sample_text == "Dr. Asha Verma"


def test_every_suggestion_carries_ui_metadata():
    report = suggest_cleaning(build_mom_docx())
    for suggestion in report.suggestions:
        assert suggestion.key and docx_parser.VALID_KEY_PATTERN.match(suggestion.key)
        assert suggestion.field_label
        assert 0 < suggestion.confidence < 1
        assert suggestion.reason
        assert suggestion.group in {
            "fixes", "lists", "tables", "blanks", "brackets",
            "empty_values", "filled_values", "detected_values",
        }


# ── 3. applying cleaning ─────────────────────────────────────────────


def test_apply_inserts_into_empty_values_and_cells_and_renders():
    for builder in (build_notice_docx, build_table_docx, build_mom_docx):
        original = builder()
        report, replacements = _accept_all(original)
        processed, results = apply_replacements_with_results(original, replacements)
        failed = [(r.placeholder_key, r.reason) for r in results if not r.matched]
        assert not failed, (builder.__name__, failed)
        # The DOCX is valid and docxtpl can parse it.
        assert zipfile.ZipFile(io.BytesIO(processed)).testzip() is None
        assert check_renderable(processed) is None
        variables = DocxTemplate(io.BytesIO(processed)).get_undeclared_template_variables()
        assert {s.key for s in report.suggestions} <= variables | {"already_there"}


def test_cleaned_notice_renders_with_sample_context():
    original = build_notice_docx()
    report, replacements = _accept_all(original)
    processed, _ = apply_replacements_with_results(original, replacements)

    document = Document(io.BytesIO(processed))
    texts = [p.text for p in document.paragraphs]
    assert "Date: {{ date }}" in texts
    assert "Name: {{ name }}\t\t\t" in texts
    assert "Department : {{ department }}" in texts
    assert "Venue {{ venue }}" in texts

    context = {key: f"VAL-{key}" for key in {s.key for s in report.suggestions}}
    context["already_there"] = "kept"
    rendered = _all_text(_render(processed, context))
    assert "Date: VAL-date" in rendered
    assert "Subject: VAL-subject" in rendered
    assert "Ref. No: VAL-ref_no" in rendered  # header
    assert "Contact: VAL-contact" in rendered  # footer
    assert "Existing kept stays untouched." in rendered
    assert "{{" not in rendered


def test_split_run_replacement_preserves_first_run_formatting():
    original = build_notice_docx()
    report = suggest_cleaning(original)
    subject = next(s for s in report.suggestions if s.key == "subject")
    processed, results = apply_replacements_with_results(
        original,
        [
            {
                "sample_text": subject.sample_text,
                "placeholder_key": "subject",
                "anchor": {
                    "paragraph_id": subject.pid,
                    "start": subject.start,
                    "end": subject.end,
                    "action": "replace",
                },
            }
        ],
    )
    assert results[0].matched
    paragraph = next(
        p for p in Document(io.BytesIO(processed)).paragraphs if p.text.startswith("Subject")
    )
    assert paragraph.text == "Subject: {{ subject }}"
    label_run = paragraph.runs[0]
    assert label_run.text == "Subject: " and label_run.bold  # label untouched
    value_run = next(r for r in paragraph.runs if "{{ subject }}" in r.text)
    assert not value_run.bold  # placeholder took the value's formatting


def test_legacy_text_replacement_keeps_formatting_and_existing_tags():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("Organised by ").bold = True
    paragraph.add_run("Dr. ")
    paragraph.add_run("Meera Rao")
    document.add_paragraph("Meera Rao signs {{ meera_rao_sig }} here")
    original = _bytes(document)
    processed, results = apply_replacements_with_results(
        original, [{"sample_text": "Dr. Meera Rao", "placeholder_key": "organiser"}]
    )
    assert results[0].occurrences == 1
    first = Document(io.BytesIO(processed)).paragraphs[0]
    assert first.text == "Organised by {{ organiser }}"
    assert first.runs[0].bold


def test_list_suggestion_becomes_loop_and_renders_items():
    original = build_mom_docx()
    report, replacements = _accept_all(original)
    processed, _ = apply_replacements_with_results(original, replacements)
    context = {key: f"[{key}]" for key in {s.key for s in report.suggestions}}
    context["agenda_items"] = ["Budget review", "Fest planning", "Any other business"]
    context["discussion_points"] = ["Point A", "Point B"]
    rendered = [p.text for p in _render(processed, context).paragraphs]
    assert rendered.count("Budget review") == 1
    assert "Any other business" in rendered
    assert "1. Point A" in rendered and "2. Point B" in rendered
    assert "ABC" not in rendered
    assert not any("{%" in text for text in rendered)


def test_textbox_twin_is_edited_in_lockstep():
    original = build_mom_docx()
    report = suggest_cleaning(original)
    hod = next(s for s in report.suggestions if s.location == "textbox")
    processed, results = apply_replacements_with_results(
        original,
        [
            {
                "sample_text": hod.sample_text,
                "placeholder_key": "hod_name",
                "anchor": {"paragraph_id": hod.pid, "start": hod.start, "end": hod.end},
            }
        ],
    )
    assert results[0].matched
    xml = zipfile.ZipFile(io.BytesIO(processed)).read("word/document.xml").decode()
    assert xml.count("{{ hod_name }}") == 2  # Choice + VML Fallback
    assert "Dr. Asha Verma" not in xml
    rendered = zipfile.ZipFile(
        io.BytesIO(_save(_render_bytes(processed, {"hod_name": "Dr. K. Iyer"})))
    ).read("word/document.xml").decode()
    assert rendered.count("Dr. K. Iyer") == 2


def _render_bytes(docx_bytes: bytes, context: dict):
    template = DocxTemplate(io.BytesIO(docx_bytes))
    template.render(context)
    return template


def _save(template) -> bytes:
    buffer = io.BytesIO()
    template.save(buffer)
    return buffer.getvalue()


def test_anchor_mismatch_and_overlaps_are_reported_not_applied():
    original = build_notice_docx()
    report = suggest_cleaning(original)
    date = next(s for s in report.suggestions if s.key == "date")
    processed, results = apply_replacements_with_results(
        original,
        [
            {"sample_text": "", "placeholder_key": "date", "anchor": {"paragraph_id": date.pid, "start": date.start, "end": date.end, "action": "insert"}},
            {"sample_text": "", "placeholder_key": "other", "anchor": {"paragraph_id": date.pid, "start": date.start, "end": date.end, "action": "insert"}},
            {"sample_text": "nope", "placeholder_key": "ghost", "anchor": {"paragraph_id": 9999, "start": 0, "end": 4}},
            {"sample_text": "x", "placeholder_key": "Bad Key"},
        ],
    )
    assert [r.reason for r in results] == [
        None,
        "overlaps_another_replacement",
        "anchor_not_found",
        "invalid_key",
    ]


def test_invalid_token_fix_suggestion():
    document = Document()
    document.add_paragraph("Dear {{Student Name}}, your marks: {{ marks }}")
    original = _bytes(document)
    report = suggest_cleaning(original)
    fixes = [s for s in report.suggestions if s.kind == "fix_invalid"]
    assert [s.key for s in fixes] == ["student_name"]
    assert not any(s.kind != "fix_invalid" and "marks" in s.sample_text for s in report.suggestions)
    _, replacements = _accept_all(original)
    processed, _ = apply_replacements_with_results(original, replacements)
    assert Document(io.BytesIO(processed)).paragraphs[0].text.startswith("Dear {{ student_name }},")
    assert check_renderable(processed) is None


# ── 4. detection ─────────────────────────────────────────────────────


def build_detection_docx() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Ref {{ref_no}}"
    document.sections[0].footer.paragraphs[0].text = "Page for {{  department  }}"
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("Meeting Details").bold = True
    split = document.add_paragraph()
    split.add_run("Held on {{ meet")
    split.add_run("ing_date }} by ")
    split.add_run("{{ chair|default('TBD') }}")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Venue"
    table.cell(0, 1).text = "{{venue}}"
    document.add_paragraph("{%p for item in attendees %}")
    document.add_paragraph("{{ item }}")
    document.add_paragraph("{%p endfor %}")
    document.add_paragraph("{{ meeting_date }} again and {{ Bad Name }}")
    return _bytes(document)


def test_detection_is_robust_across_runs_tables_headers_loops():
    result = docx_parser.detect_placeholders(build_detection_docx())
    keys = [field.key for field in result.valid_fields]
    assert keys == ["meeting_date", "chair", "venue", "attendees", "ref_no", "department"]
    by_key = {field.key: field for field in result.valid_fields}
    assert by_key["attendees"].type_hint == "list"
    assert by_key["meeting_date"].section == "Meeting Details"
    assert by_key["ref_no"].section == "Header"
    assert by_key["department"].location == "footer"
    assert "item" not in keys
    assert [d.key for d in result.duplicates] == ["meeting_date"]
    assert [i.raw for i in result.invalid_names] == ["Bad Name"]
    assert result.invalid_names[0].suggested_key == "bad_name"
    assert result.parse_warning is not None  # {{ Bad Name }} blocks docxtpl


# ── 5. endpoints ─────────────────────────────────────────────────────


def _signup_and_login(email: str, full_name: str) -> str:
    client.post(
        "/api/v1/auth/signup",
        json={"full_name": full_name, "email": email, "password": "strong-password"},
    )
    login = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "strong-password"}
    )
    return login.json()["access_token"]


@pytest.fixture(scope="module")
def owner_token():
    return _signup_and_login("pipeline_owner@example.com", "Pipeline Owner")


@pytest.fixture(scope="module")
def other_token():
    return _signup_and_login("pipeline_other@example.com", "Pipeline Other")


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _upload(token: str, name: str, docx_bytes: bytes) -> dict:
    response = client.post(
        "/api/v1/templates/upload",
        files={
            "file": (
                "pipeline.docx",
                io.BytesIO(docx_bytes),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"name": name, "category": "mom", "visibility": "private"},
        headers=_auth(token),
    )
    assert response.status_code == 201, response.text
    return response.json()


def _accept_payload(suggestions: list[dict]) -> list[dict]:
    return [
        {
            "sample_text": s["sample_text"],
            "placeholder_key": s["placeholder_key"],
            "field_label": s["field_label"],
            "field_type": s["field_type"],
            "section": s["section"],
            "example_value": s["example_value"],
            "validation_rule": s["validation_rule"],
            "is_required": s["is_required"],
            "anchor": s["anchor"],
        }
        for s in suggestions
    ]


def test_clean_suggestions_endpoint_contract(owner_token, other_token):
    template = _upload(owner_token, "Pipeline Suggestions", build_mom_docx())
    url = f"/api/v1/templates/{template['id']}/clean-suggestions"

    response = client.get(url, headers=_auth(owner_token))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["suggestion_count"] == len(body["suggestions"]) > 5
    assert body["high_confidence_threshold"] == 0.8
    assert body["high_confidence_count"] == sum(s["high_confidence"] for s in body["suggestions"])
    first = body["suggestions"][0]
    for field in ("id", "kind", "group", "placeholder_key", "field_label", "field_type",
                  "confidence", "reason", "anchor", "context", "sample_text", "location"):
        assert field in first
    assert client.get(url, headers=_auth(other_token)).status_code == 403
    assert client.get("/api/v1/templates/999999/clean-suggestions", headers=_auth(owner_token)).status_code == 404


def test_preview_then_clean_then_detect_flow(owner_token):
    template = _upload(owner_token, "Pipeline Full Flow", build_notice_docx())
    tid = template["id"]
    suggestions = client.get(
        f"/api/v1/templates/{tid}/clean-suggestions", headers=_auth(owner_token)
    ).json()["suggestions"]
    accepted = _accept_payload([s for s in suggestions if s["high_confidence"]])
    assert accepted

    # Preview: nothing persisted, confirm not needed.
    preview = client.post(
        f"/api/v1/templates/{tid}/clean-preview",
        headers=_auth(owner_token),
        json={"replacements": accepted},
    )
    assert preview.status_code == 200, preview.text
    preview_body = preview.json()
    assert any("{{ date }}" in seg["text"] for seg in preview_body["segments"])
    assert preview_body["warnings"]["render_error"] is None
    detail = client.get(f"/api/v1/templates/{tid}", headers=_auth(owner_token)).json()
    assert detail["processed_file_path"] is None

    # Clean: fields carry inferred metadata.
    cleaned = client.post(
        f"/api/v1/templates/{tid}/clean",
        headers=_auth(owner_token),
        json={"replacements": accepted, "confirm": True},
    )
    assert cleaned.status_code == 200, cleaned.text
    body = cleaned.json()
    assert all(r["matched"] for r in body["results"]), body["results"]
    fields = {f["field_name"]: f for f in body["created_fields"]}
    assert fields["date"]["field_type"] == "date"
    assert fields["date"]["example_value"] is None  # inserted, nothing to keep
    assert fields["subject"]["example_value"] == "Annual Sports Meet"
    assert fields["signature"]["field_type"] == "signature"
    assert fields["ref_no"]["section"] == "Header"
    assert "date" in body["detected_keys"]
    assert body["warnings"]["render_error"] is None

    # Detection now scans the processed copy and keeps the cleaned fields.
    detected = client.post(
        f"/api/v1/templates/{tid}/detect-placeholders", headers=_auth(owner_token)
    )
    assert detected.status_code == 200, detected.text
    detection = detected.json()
    assert detection["source_file"] == "processed"
    names = [f["field_name"] for f in detection["detected_fields"]]
    assert set(fields) <= set(names)
    assert "already_there" in detection["added_fields"]  # pre-existing tag
    assert detection["stale_fields"] == []

    # Force re-detect rebuilds in document order and keeps example values.
    forced = client.post(
        f"/api/v1/templates/{tid}/detect-placeholders?force=true", headers=_auth(owner_token)
    ).json()
    forced_fields = {f["field_name"]: f for f in forced["detected_fields"]}
    assert forced_fields["subject"]["example_value"] == "Annual Sports Meet"
    assert [f["display_order"] for f in forced["detected_fields"]] == list(range(len(forced_fields)))


def test_reclean_reports_stale_fields(owner_token):
    template = _upload(owner_token, "Pipeline Stale", build_notice_docx())
    tid = template["id"]
    headers = _auth(owner_token)
    first = client.post(
        f"/api/v1/templates/{tid}/clean",
        headers=headers,
        json={"replacements": [
            {"sample_text": "Annual Sports Meet", "placeholder_key": "subject"},
            {"sample_text": "hod.cse@univ.edu", "placeholder_key": "contact_email"},
        ], "confirm": True},
    ).json()
    assert {f["field_name"] for f in first["created_fields"]} == {"subject", "contact_email"}
    assert first["created_fields"][1]["validation_rule"] == "email"
    second = client.post(
        f"/api/v1/templates/{tid}/clean",
        headers=headers,
        json={"replacements": [{"sample_text": "Annual Sports Meet", "placeholder_key": "subject"}], "confirm": True},
    ).json()
    assert second["warnings"]["stale_fields"] == ["contact_email"]


def test_fix_placeholders_endpoint(owner_token, other_token):
    document = Document()
    document.add_paragraph("Dear {{Student Name}}, see {{ marks }}")
    template = _upload(owner_token, "Pipeline Fix", _bytes(document))
    tid = template["id"]
    detect = client.post(
        f"/api/v1/templates/{tid}/detect-placeholders", headers=_auth(owner_token)
    ).json()
    assert detect["warnings"]["invalid_names"][0]["suggested_key"] == "student_name"

    url = f"/api/v1/templates/{tid}/fix-placeholders"
    payload = {"fixes": [{"raw": "Student Name", "key": "student_name"}]}
    assert client.post(url, headers=_auth(other_token), json=payload).status_code == 403
    response = client.post(url, headers=_auth(owner_token), json=payload)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["warnings"]["invalid_names"] == []
    assert body["warnings"]["parse_error"] is None
    assert body["added_fields"] == ["student_name"]
    assert body["source_file"] == "processed"
    missing = client.post(url, headers=_auth(owner_token), json={"fixes": [{"raw": "Nope", "key": "nope"}]})
    assert missing.status_code == 422


def test_clean_request_contract_is_backward_compatible(owner_token):
    template = _upload(owner_token, "Pipeline Compat", build_notice_docx())
    url = f"/api/v1/templates/{template['id']}/clean"
    headers = _auth(owner_token)
    # Old shape (no anchor) still works; empty sample_text without anchor is rejected.
    ok = client.post(url, headers=headers, json={"replacements": [{"sample_text": "Sports Meet", "placeholder_key": "event"}], "confirm": True})
    assert ok.status_code == 200
    bad = client.post(url, headers=headers, json={"replacements": [{"sample_text": "", "placeholder_key": "event"}], "confirm": True})
    assert bad.status_code == 422
    bad_anchor = client.post(url, headers=headers, json={"replacements": [{"placeholder_key": "event", "anchor": {"paragraph_id": 1, "start": 5, "end": 2}}], "confirm": True})
    assert bad_anchor.status_code == 422


# ── 6. field editor robustness ───────────────────────────────────────


def _make_template_row(name: str) -> int:
    with TestSession() as session:
        owner = session.query(User).filter_by(email="pipeline_owner@example.com").one()
        template = Template(name=name, category="mom", visibility="private", uploaded_by=owner.id)
        session.add(template)
        session.commit()
        return template.id


def _fields(token: str, template_id: int) -> list[dict]:
    return client.get(f"/api/v1/templates/{template_id}/fields", headers=_auth(token)).json()


def test_create_field_validates_and_infers(owner_token):
    tid = _make_template_row("Editor Infer")
    url = f"/api/v1/templates/{tid}/fields"
    headers = _auth(owner_token)
    assert client.post(url, headers=headers, json={"field_name": "Bad Name"}).status_code == 422
    assert client.post(url, headers=headers, json={"field_name": "ok", "validation_rule": "regex:.*"}).status_code == 422
    created = client.post(url, headers=headers, json={"field_name": "meeting_date"})
    assert created.status_code == 201
    assert created.json()["field_type"] == "date"
    assert created.json()["field_label"] == "Meeting Date"
    explicit = client.post(url, headers=headers, json={"field_name": "notes_text", "field_type": "text", "field_label": "Notes"})
    assert explicit.json()["field_type"] == "text"
    assert explicit.json()["field_label"] == "Notes"
    email = client.post(url, headers=headers, json={"field_name": "contact_email"}).json()
    assert email["validation_rule"] == "email"
    assert client.post(url, headers=headers, json={"field_name": "meeting_date"}).status_code == 409


def test_create_and_update_respect_explicit_position(owner_token):
    tid = _make_template_row("Editor Positions")
    url = f"/api/v1/templates/{tid}/fields"
    headers = _auth(owner_token)
    for name in ("alpha", "beta", "gamma"):
        client.post(url, headers=headers, json={"field_name": name})
    client.post(url, headers=headers, json={"field_name": "first", "display_order": 0})
    client.post(url, headers=headers, json={"field_name": "middle", "display_order": 2})
    names = [f["field_name"] for f in _fields(owner_token, tid)]
    assert names == ["first", "alpha", "middle", "beta", "gamma"]
    gamma = next(f for f in _fields(owner_token, tid) if f["field_name"] == "gamma")
    moved = client.patch(f"{url}/{gamma['id']}", headers=headers, json={"display_order": 1})
    assert moved.status_code == 200
    fields = _fields(owner_token, tid)
    assert [f["field_name"] for f in fields] == ["first", "gamma", "alpha", "middle", "beta"]
    assert [f["display_order"] for f in fields] == [0, 1, 2, 3, 4]
    assert client.patch(f"{url}/{gamma['id']}", headers=headers, json={"validation_rule": "max:abc"}).status_code == 422
    blank = client.patch(f"{url}/{gamma['id']}", headers=headers, json={"field_label": "   ", "validation_rule": " MIN:2 | max:5 "})
    assert blank.json()["field_label"] is None
    assert blank.json()["validation_rule"] == "min:2|max:5"


def test_sync_is_idempotent_and_reorder_closes_gaps(owner_token):
    tid = _make_template_row("Editor Sync")
    headers = _auth(owner_token)
    url = f"/api/v1/templates/{tid}/fields"
    payload = {"fields": [
        {"field_name": "agenda_items"},
        {"field_name": "meeting_date", "field_label": "When"},
        {"field_name": "remarks", "field_type": "textarea", "is_required": False},
    ]}
    first = client.put(url, headers=headers, json=payload).json()
    assert [f["field_type"] for f in first] == ["list", "date", "textarea"]
    assert first[1]["field_label"] == "When"
    again = {"fields": [{**item, "id": row["id"]} for item, row in zip(payload["fields"], first)]}
    second = client.put(url, headers=headers, json=again).json()
    third = client.put(url, headers=headers, json=again).json()
    strip = lambda rows: [{k: v for k, v in r.items() if k != "updated_at"} for r in rows]  # noqa: E731
    assert strip(second) == strip(third)
    assert [f["id"] for f in second] == [f["id"] for f in first]

    bad_type = client.put(url, headers=headers, json={"fields": [{"field_name": "x", "field_type": "dropdown"}]})
    assert bad_type.status_code == 422
    dup = client.put(url, headers=headers, json={"fields": [{"field_name": "x"}, {"field_name": "x"}]})
    assert dup.status_code == 422

    # Delete in the middle, then reorder: orders stay contiguous.
    client.delete(f"{url}/{first[1]['id']}", headers=headers)
    remaining = _fields(owner_token, tid)
    assert [f["display_order"] for f in remaining] == [0, 1]
    reordered = client.put(f"{url}/reorder", headers=headers, json={"ordered_ids": [remaining[1]["id"], remaining[0]["id"]]})
    assert [f["field_name"] for f in reordered.json()] == ["remarks", "agenda_items"]
    gap = client.put(f"{url}/reorder", headers=headers, json={"ordered_ids": [remaining[0]["id"]]})
    assert gap.status_code == 422


def test_example_values_never_overflow(owner_token):
    long_value = "x" * 300
    document = Document()
    document.add_paragraph(f"Remarks: {long_value}")
    template = _upload(owner_token, "Pipeline Long", _bytes(document))
    response = client.post(
        f"/api/v1/templates/{template['id']}/clean",
        headers=_auth(owner_token),
        json={"replacements": [{"sample_text": long_value, "placeholder_key": "remarks"}], "confirm": True},
    )
    assert response.status_code == 200, response.text
    assert len(response.json()["created_fields"][0]["example_value"]) == 255


def test_json_list_example_for_list_suggestions():
    report = suggest_cleaning(build_mom_docx())
    agenda = next(s for s in report.suggestions if s.key == "agenda_items")
    assert agenda.example_value is None  # sample items were filler ("ABC")
    json.dumps([s.__dict__.get("key") for s in report.suggestions])
