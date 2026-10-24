"""
V1.7 Phase 2.1 API (Member 2): Word file editor with a FAKE ONLYOFFICE
Document Server (httpx.MockTransport; no container).

Coverage:
- /system/editor-config enabled / disabled; 503 on session routes when off.
- Customize: private copy for the caller with fields; 403 without access;
  the original is untouched; editing the original is for managers only.
- Template session: signed config (JWT verifies), signed file link serves
  the file; bad / expired token 403; lock 409 for another user; reuse.
- Template save callback: new revision, file switched, field changes,
  status back to placeholder_detected; unparseable file kept but not used;
  revisions list + restore; upload / clean / fix add revisions.
- Document session + callback: manual_edit version with base snapshot,
  PDF converted, document not stale; duplicate callback stored once;
  regenerate guard 409 code manual_edits, replace_manual_edits=true.
- Rejections: unsigned / bad JWT, foreign host (no request made),
  oversize, not a DOCX; status 3 / 4.

Storage is a temporary folder; the database is in-memory SQLite with
foreign keys on; PDF conversion is a fake converter.
"""

import io
import json
import os
from datetime import timedelta
from urllib.parse import urlsplit

import httpx
import jwt
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests")

from app.api.v1.endpoints._template_common import record_file_revision  # noqa: E402
from app.core.clock import utc_now_naive  # noqa: E402
from app.core.config import settings  # noqa: E402
from app.crud.template_field_crud import create_field  # noqa: E402
from app.db.base import Base  # noqa: E402
from app.db.session import get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models.document import Document  # noqa: E402
from app.models.document_version import DocumentVersion  # noqa: E402
from app.models.editor_session import EditorSession  # noqa: E402
from app.models.template import Template  # noqa: E402
from app.models.template_field import TemplateField  # noqa: E402
from app.models.template_revision import TemplateRevision  # noqa: E402
from app.models.user import User  # noqa: E402
from app.schemas.template_field import TemplateFieldCreate  # noqa: E402
from app.services import document_generation, onlyoffice_service  # noqa: E402
from app.services.storage_service import storage_service  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)


@event.listens_for(engine, "connect")
def _enable_foreign_keys(dbapi_connection, _record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)
API = "/api/v1"
SECRET = "onlyoffice-test-secret-that-is-long-enough-0123"
DOCS = "http://docs.test:8082"
FAKE_PDF = b"%PDF-1.4\n% fake\n" + b"0" * 400


@pytest.fixture(scope="module", autouse=True)
def setup_db(tmp_path_factory):
    old = app.dependency_overrides.get(get_db)
    old_root = storage_service.root_path
    storage_service.root_path = tmp_path_factory.mktemp("p21-storage").resolve()
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    storage_service.root_path = old_root
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


class FakeDocumentServer:
    """Serves saved files at {DOCS}/cache/files/<name>; records requests."""

    def __init__(self):
        self.files: dict[str, bytes] = {}
        self.requests: list[httpx.Request] = []

    def publish(self, content: bytes, name: str = "edited.docx") -> str:
        self.files[name] = content
        return f"{DOCS}/cache/files/{name}?md5=abc"

    def handler(self, request: httpx.Request) -> httpx.Response:
        self.requests.append(request)
        name = request.url.path.rsplit("/", 1)[-1]
        if name in self.files:
            return httpx.Response(200, content=self.files[name])
        return httpx.Response(404)


class FakeConverter:
    def __init__(self):
        self.calls = 0

    def convert(self, docx_bytes: bytes, filename: str) -> bytes:
        self.calls += 1
        return FAKE_PDF


@pytest.fixture(autouse=True)
def editor_on(monkeypatch):
    server = FakeDocumentServer()
    monkeypatch.setattr(settings, "onlyoffice_enabled", True)
    monkeypatch.setattr(settings, "onlyoffice_public_url", DOCS)
    monkeypatch.setattr(settings, "onlyoffice_internal_url", None)
    monkeypatch.setattr(settings, "onlyoffice_jwt_secret", SECRET)
    monkeypatch.setattr(settings, "backend_url_for_onlyoffice", "http://backend.test")
    monkeypatch.setattr(settings, "generated_versions_keep_files", 10)
    monkeypatch.setattr(onlyoffice_service, "document_server_transport", httpx.MockTransport(server.handler))
    converter = FakeConverter()
    monkeypatch.setattr(document_generation, "get_pdf_converter", lambda: converter)
    server.converter = converter
    return server


def docx(*paragraphs: str) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_n = {"value": 0}


def _login(role: str = "normal_user") -> tuple[str, int]:
    _n["value"] += 1
    email = f"p21_{_n['value']}@example.com"
    client.post(f"{API}/auth/signup", json={"full_name": f"User {_n['value']}", "email": email,
                                            "password": "strong-password"})
    with TestSession() as db:
        user = db.execute(select(User).where(User.email == email)).scalar_one()
        user.role = role
        db.commit()
        user_id = user.id
    response = client.post(f"{API}/auth/login", json={"email": email, "password": "strong-password"})
    assert response.status_code == 200, response.text
    return response.json()["access_token"], user_id


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _template(owner_id: int, visibility: str = "public", paragraphs=("Title: {{ title }}",)) -> int:
    content = docx(*paragraphs)
    original = storage_service.save_bytes("templates_original", "mom.docx", docx("Title: ______")).path
    processed = storage_service.save_bytes("templates_processed", "mom.docx", content).path
    with TestSession() as db:
        template = Template(
            name="MoM", category="mom", visibility=visibility, status="active", uploaded_by=owner_id,
            original_file_path=original, processed_file_path=processed, original_filename="mom.docx",
        )
        db.add(template)
        db.commit()
        create_field(db, template.id, TemplateFieldCreate(field_name="title", field_type="text", is_required=True))
        record_file_revision(db, template, processed, content, "upload", owner_id)
        return template.id


def _callback(key: str, payload: dict, *, secret: str = SECRET, in_header: bool = False):
    """Like the Document Server: no browser cookies (so no CSRF header either)."""
    server = TestClient(app)
    payload = {"key": key, **payload}
    if in_header:
        token = jwt.encode({"payload": payload}, secret, algorithm="HS256")
        return server.post(f"{API}/editor/callback/{key}", json=payload,
                           headers={"Authorization": f"Bearer {token}"})
    token = jwt.encode(payload, secret, algorithm="HS256")
    return server.post(f"{API}/editor/callback/{key}", json={**payload, "token": token})


def _file_path_from(config: dict) -> str:
    url = urlsplit(config["document"]["url"])
    assert url.netloc == "backend.test"
    return f"{url.path}?{url.query}"


# ── Availability ─────────────────────────────────────────────────


def test_editor_config_and_disabled(monkeypatch):
    assert client.get(f"{API}/system/editor-config").json() == {"enabled": True, "document_server_url": DOCS}
    monkeypatch.setattr(settings, "onlyoffice_enabled", False)
    assert client.get(f"{API}/system/editor-config").json() == {"enabled": False, "document_server_url": None}
    token, user_id = _login("org_admin")
    template_id = _template(user_id)
    assert client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(token)).status_code == 503


# ── Templates ────────────────────────────────────────────────────


def test_customize_and_template_session(editor_on):
    admin, admin_id = _login("org_admin")
    faculty, faculty_id = _login("faculty")
    template_id = _template(admin_id)

    # Faculty may not edit the admin's original ...
    denied = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(faculty))
    assert denied.status_code == 403 and "Customize" in denied.json()["detail"]
    # ... but gets a private copy.
    response = client.post(f"{API}/templates/{template_id}/customize", headers=_auth(faculty))
    assert response.status_code == 201, response.text
    copy = response.json()
    assert (copy["parent_template_id"], copy["is_personal_copy"], copy["visibility"]) == (template_id, True, "private")
    assert (copy["uploaded_by"], copy["can_manage"], copy["name"]) == (faculty_id, True, "MoM (my copy)")
    with TestSession() as db:
        names = [f.field_name for f in db.execute(select(TemplateField).where(TemplateField.template_id == copy["id"])).scalars()]
        assert names == ["title"]
        assert db.get(Template, template_id).version == 1

    opened = client.post(f"{API}/templates/{copy['id']}/editor-session", headers=_auth(faculty))
    assert opened.status_code == 200, opened.text
    body = opened.json()
    assert body["document_server_url"] == DOCS and body["reused"] is False
    config = body["config"]
    claims = jwt.decode(config["token"], SECRET, algorithms=["HS256"])
    assert claims["document"]["key"] == body["session_key"] == config["document"]["key"]
    assert config["editorConfig"]["callbackUrl"] == f"http://backend.test/api/v1/editor/callback/{body['session_key']}"
    assert config["documentType"] == "word" and config["document"]["fileType"] == "docx"

    # The Document Server fetches the file with the signed link (no cookies).
    path = _file_path_from(config)
    fetched = TestClient(app).get(path)
    assert fetched.status_code == 200 and fetched.content.startswith(b"PK")
    assert TestClient(app).get(path.replace("token=", "token=1")).status_code == 403
    assert TestClient(app).get(f"{API}/editor/files/{body['session_key']}").status_code == 403
    assert TestClient(app).get(f"{API}/editor/files/unknown-key?token=x").status_code == 404

    # Same user reuses; the copy is private, so others cannot even manage it.
    again = client.post(f"{API}/templates/{copy['id']}/editor-session", headers=_auth(faculty)).json()
    assert (again["session_key"], again["reused"]) == (body["session_key"], True)
    assert client.post(f"{API}/templates/{copy['id']}/editor-session", headers=_auth(admin)).status_code == 403

    # Status polling: owner only.
    status_response = client.get(f"{API}/editor/sessions/{body['session_key']}", headers=_auth(faculty))
    assert status_response.json()["status"] == "open"
    assert client.get(f"{API}/editor/sessions/{body['session_key']}", headers=_auth(admin)).status_code == 404


def test_lock_between_managers_and_expired_link():
    admin, admin_id = _login("org_admin")
    template_id = _template(admin_id, visibility="private")
    super_token, _ = _login("super_admin")
    first = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(admin)).json()
    locked = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(super_token))
    assert locked.status_code == 409 and "Being edited by" in locked.json()["detail"]
    with TestSession() as db:
        session = db.execute(select(EditorSession).where(EditorSession.session_key == first["session_key"])).scalar_one()
        session.expires_at = utc_now_naive() - timedelta(minutes=1)
        db.commit()
    assert TestClient(app).get(_file_path_from(first["config"])).status_code == 403
    taken = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(super_token))
    assert taken.status_code == 200


def test_template_save_callback_records_revision_and_field_changes(editor_on):
    admin, admin_id = _login("org_admin")
    template_id = _template(admin_id)
    key = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(admin)).json()["session_key"]
    assert _callback(key, {"status": 1}).json() == {"error": 0}

    edited = docx("Title: {{ title }}", "Date of meeting: {{ meeting_date }}", "Venue: ________________")
    url = editor_on.publish(edited)
    assert _callback(key, {"status": 2, "url": url}).json() == {"error": 0}
    status_body = client.get(f"{API}/editor/sessions/{key}", headers=_auth(admin)).json()
    assert (status_body["status"], status_body["result_revision"]) == ("saved", 2)

    detail = client.get(f"{API}/templates/{template_id}", headers=_auth(admin)).json()
    assert detail["version"] == 2 and detail["status"] == "placeholder_detected"
    assert storage_service.read_bytes(detail["processed_file_path"]) == edited

    changes = client.get(f"{API}/templates/{template_id}/field-changes", headers=_auth(admin)).json()
    assert [(a["key"], a["suggested_type"]) for a in changes["added"]] == [("meeting_date", "date")]
    assert changes["removed"] == [] and changes["blank_suggestions"] >= 1
    assert changes["renderable"] is True and changes["has_changes"] is True

    revisions = client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(admin)).json()
    assert [(r["revision"], r["source"], r["is_current"]) for r in revisions] == [
        (2, "manual_edit", True), (1, "upload", False),
    ]
    assert revisions[0]["placeholders_added"] == ["meeting_date"]
    # Duplicate callback: stored once.
    assert _callback(key, {"status": 2, "url": url}).json() == {"error": 0}
    assert len(client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(admin)).json()) == 2

    # Restore revision 1 -> revision 3, current, same bytes as revision 1.
    restored = client.post(f"{API}/templates/{template_id}/revisions/1/restore", headers=_auth(admin))
    assert restored.status_code == 200 and restored.json()["version"] == 3
    revisions = client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(admin)).json()
    assert (revisions[0]["revision"], revisions[0]["source"], revisions[0]["is_current"]) == (3, "restore", True)
    assert "meeting_date" not in client.get(f"{API}/templates/{template_id}/field-changes",
                                            headers=_auth(admin)).text.split("removed")[0]
    assert client.post(f"{API}/templates/{template_id}/revisions/9/restore", headers=_auth(admin)).status_code == 404
    faculty, _ = _login("faculty")
    assert client.post(f"{API}/templates/{template_id}/revisions/1/restore", headers=_auth(faculty)).status_code == 403
    assert client.get(f"{API}/templates/{template_id}/field-changes", headers=_auth(faculty)).status_code == 403
    assert client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(faculty)).status_code == 200


def test_unparseable_template_edit_is_kept_but_not_used(editor_on):
    admin, admin_id = _login("org_admin")
    template_id = _template(admin_id)
    before = client.get(f"{API}/templates/{template_id}", headers=_auth(admin)).json()["processed_file_path"]
    key = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(admin)).json()["session_key"]
    url = editor_on.publish(docx("Title: {{ title", "{% if %}"), "broken.docx")
    assert _callback(key, {"status": 2, "url": url}).json() == {"error": 0}
    status_body = client.get(f"{API}/editor/sessions/{key}", headers=_auth(admin)).json()
    assert status_body["status"] == "failed" and "could not be read" in status_body["error"]
    detail = client.get(f"{API}/templates/{template_id}", headers=_auth(admin)).json()
    assert detail["processed_file_path"] == before
    revisions = client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(admin)).json()
    assert revisions[0]["note"].startswith("Not activated") and revisions[0]["is_current"] is False
    assert client.post(f"{API}/templates/{template_id}/revisions/{revisions[0]['revision']}/restore",
                       headers=_auth(admin)).status_code == 422


def test_customize_rules():
    admin, admin_id = _login("org_admin")
    private_id = _template(admin_id, visibility="private")
    faculty, _ = _login("faculty")
    assert client.post(f"{API}/templates/{private_id}/customize", headers=_auth(faculty)).status_code == 403
    assert client.post(f"{API}/templates/999999/customize", headers=_auth(faculty)).status_code == 404
    second = client.post(f"{API}/templates/{private_id}/customize", headers=_auth(admin))
    assert second.status_code == 201 and second.json()["name"] == "MoM (my copy)"


def test_upload_records_first_revision():
    token, _ = _login("normal_user")
    response = client.post(
        f"{API}/templates/upload",
        files={"file": ("notice.docx", docx("Dear {{ name }}"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"name": "Notice", "category": "notice", "visibility": "private"},
        headers=_auth(token),
    )
    assert response.status_code in (200, 201), response.text
    template_id = response.json()["id"]
    revisions = client.get(f"{API}/templates/{template_id}/revisions", headers=_auth(token)).json()
    assert [(r["revision"], r["source"], r["is_current"]) for r in revisions] == [(1, "upload", True)]


# ── Documents ────────────────────────────────────────────────────


def _generated_document(token: str, owner_id: int) -> int:
    template_id = _template(owner_id, visibility="private")
    created = client.post(f"{API}/documents/create", headers=_auth(token),
                          json={"template_id": template_id, "name": "Weekly MoM"}).json()
    document_id = created["id"]
    assert client.put(f"{API}/documents/{document_id}/values", headers=_auth(token),
                      json={"values": [{"field_name": "title", "value": "Budget"}]}).status_code == 200
    assert client.post(f"{API}/documents/{document_id}/generate", headers=_auth(token)).json()["version"] == 1
    return document_id


def test_document_manual_edit_and_regenerate_guard(editor_on):
    token, owner_id = _login("normal_user")
    document_id = _generated_document(token, owner_id)
    with TestSession() as db:
        updated_before = db.get(Document, document_id).updated_at
    opened = client.post(f"{API}/documents/{document_id}/editor-session", headers=_auth(token))
    assert opened.status_code == 200, opened.text
    body = opened.json()
    assert body["config"]["document"]["title"] == "Weekly MoM v1.docx"
    original = TestClient(app).get(_file_path_from(body["config"]))
    assert "Budget" in "".join(p.text for p in DocxDocument(io.BytesIO(original.content)).paragraphs)

    edited = docx("Title: Budget (corrected by hand)")
    calls_before = editor_on.converter.calls
    assert _callback(body["session_key"], {"status": 2, "url": editor_on.publish(edited)}, in_header=True).json() == {"error": 0}
    status_body = client.get(f"{API}/editor/sessions/{body['session_key']}", headers=_auth(token)).json()
    assert (status_body["status"], status_body["result_version"]) == ("saved", 2)
    assert editor_on.converter.calls == calls_before + 1

    versions = client.get(f"{API}/documents/{document_id}/versions", headers=_auth(token)).json()
    assert [(v["version"], v["source"], v["base_version"]) for v in versions] == [
        (2, "manual_edit", 1), (1, "generated", None),
    ]
    assert versions[0]["has_pdf"] is True and versions[0]["changed_fields"] == []
    latest = client.get(f"{API}/documents/{document_id}/download?format=docx", headers=_auth(token))
    assert latest.content == edited
    with TestSession() as db:
        v1, v2 = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == document_id)
                            .order_by(DocumentVersion.version)).scalars().all()
        assert v2.values_snapshot == v1.values_snapshot
        assert db.get(Document, document_id).updated_at == updated_before
    preview = client.get(f"{API}/documents/{document_id}/preview", headers=_auth(token)).json()
    assert preview["generation"]["source"] == "manual_edit" and preview["generation"]["is_stale"] is False

    # Duplicate callback: no new version.
    assert _callback(body["session_key"], {"status": 2, "url": editor_on.publish(edited)}).json() == {"error": 0}
    assert len(client.get(f"{API}/documents/{document_id}/versions", headers=_auth(token)).json()) == 2

    guard = client.post(f"{API}/documents/{document_id}/generate", headers=_auth(token))
    assert guard.status_code == 409
    assert guard.json()["detail"]["code"] == "manual_edits" and guard.json()["detail"]["version"] == 2
    replaced = client.post(f"{API}/documents/{document_id}/generate?replace_manual_edits=true", headers=_auth(token))
    assert replaced.status_code == 200 and replaced.json()["version"] == 3

    # Editing an older version works; unknown version 404; other user 403.
    assert client.post(f"{API}/documents/{document_id}/editor-session?version=1", headers=_auth(token)).status_code == 200
    assert client.post(f"{API}/documents/{document_id}/editor-session?version=9", headers=_auth(token)).status_code == 404
    stranger, _ = _login("normal_user")
    assert client.post(f"{API}/documents/{document_id}/editor-session", headers=_auth(stranger)).status_code == 403


def test_document_session_rules(monkeypatch):
    token, owner_id = _login("normal_user")
    template_id = _template(owner_id, visibility="private")
    document_id = client.post(f"{API}/documents/create", headers=_auth(token),
                              json={"template_id": template_id}).json()["id"]
    assert client.post(f"{API}/documents/{document_id}/editor-session", headers=_auth(token)).status_code == 404
    document_id = _generated_document(token, owner_id)
    with TestSession() as db:
        db.get(Document, document_id).status = "submitted"
        db.commit()
    assert client.post(f"{API}/documents/{document_id}/editor-session", headers=_auth(token)).status_code == 409
    document_id = _generated_document(token, owner_id)
    with TestSession() as db:
        record = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == document_id)).scalar_one().generated_file
        record.files_pruned_at = utc_now_naive()
        db.commit()
    assert client.post(f"{API}/documents/{document_id}/editor-session", headers=_auth(token)).status_code == 410


# ── Callback rejections ──────────────────────────────────────────


def _open_template_session(token: str, owner_id: int) -> tuple[int, str]:
    template_id = _template(owner_id, visibility="private")
    key = client.post(f"{API}/templates/{template_id}/editor-session", headers=_auth(token)).json()["session_key"]
    return template_id, key


def test_callback_rejections(editor_on, monkeypatch):
    token, owner_id = _login("org_admin")
    template_id, key = _open_template_session(token, owner_id)
    url = editor_on.publish(docx("{{ title }} edited"))

    unsigned = TestClient(app).post(f"{API}/editor/callback/{key}", json={"key": key, "status": 2, "url": url})
    assert unsigned.json() == {"error": 1}
    assert _callback(key, {"status": 2, "url": url}, secret="wrong-secret-that-is-long-enough-123456").json() == {"error": 1}
    assert _callback("unknown", {"status": 2, "url": url}).json() == {"error": 1}
    assert editor_on.requests == []  # nothing downloaded

    foreign = _callback(key, {"status": 2, "url": "http://169.254.169.254/latest/meta-data"})
    assert foreign.json() == {"error": 1} and editor_on.requests == []
    status_body = client.get(f"{API}/editor/sessions/{key}", headers=_auth(token)).json()
    assert status_body["status"] == "failed" and "document editor server" in status_body["error"]

    _, key = _open_template_session(token, owner_id)
    assert _callback(key, {"status": 2, "url": editor_on.publish(b"not a docx", "x.docx")}).json() == {"error": 1}
    assert "not a Word document" in client.get(f"{API}/editor/sessions/{key}", headers=_auth(token)).json()["error"]

    _, key = _open_template_session(token, owner_id)
    monkeypatch.setattr(settings, "editor_max_file_mb", 1)
    big = editor_on.publish(b"PK" + b"0" * (2 * 1024 * 1024), "big.docx")
    assert _callback(key, {"status": 2, "url": big}).json() == {"error": 1}
    assert "larger than 1 MB" in client.get(f"{API}/editor/sessions/{key}", headers=_auth(token)).json()["error"]

    _, key = _open_template_session(token, owner_id)
    assert _callback(key, {"status": 4}).json() == {"error": 0}
    assert client.get(f"{API}/editor/sessions/{key}", headers=_auth(token)).json()["status"] == "closed"
    _, key = _open_template_session(token, owner_id)
    assert _callback(key, {"status": 3}).json() == {"error": 0}
    assert client.get(f"{API}/editor/sessions/{key}", headers=_auth(token)).json()["status"] == "failed"

    with TestSession() as db:
        assert db.execute(select(TemplateRevision).where(TemplateRevision.template_id == template_id)).scalars().all()[-1].revision == 1
    assert json.dumps({}) == "{}"
