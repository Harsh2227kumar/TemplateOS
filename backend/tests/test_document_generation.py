"""
V1.6 document generation: DOCX rendering, PDF conversion, download, cleanup.

Coverage:
- generation_service: lists as Python lists (docxtpl {%p for %} loops),
  ISO dates formatted, empty / optional values blank (never "None"),
  escaping, newlines, template errors (undefined variable, syntax error).
- POST /generate: owner-only (403 / 404), required-field 422 with
  {field_name, error}, template error 422, status -> generated, generated
  info on GET /documents/{id} and the list, staleness after a later save.
- PDF: mocked Gotenberg success / HTTP failure / timeout, not configured,
  POST /generate-pdf retry, LibreOffice converter (mocked subprocess) and
  converter selection.
- GET /download: headers, filename from the draft name, 404 / 409 cases.
- Regeneration replaces files; document / template delete removes rows and
  stored files.
- Migration up / down / up.
- End to end: the real Minutes-of-Meeting sample cleaned through the clean
  API, then generated (python-docx fixture if the sample is unavailable).

Storage is redirected to a temporary folder; the database is in-memory
SQLite with foreign keys enforced (so ON DELETE CASCADE is exercised).
"""

import html
import io
import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import httpx
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, inspect, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault(
    "JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests"
)

from app.core.config import settings
from app.crud.document_crud import delete_document
from app.crud.template_field_crud import create_field
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.document import Document
from app.models.generated_file import GeneratedFile
from app.models.template import Template
from app.models.template_field import TemplateField
from app.models.user import User
from app.schemas.template_field import TemplateFieldCreate
from app.services import pdf_service
from app.services.document_generation import download_filename
from app.services.generation_service import (
    ListValue,
    TemplateRenderError,
    build_context,
    format_date_value,
    render_docx,
)
from app.services.storage_service import storage_service

BACKEND_DIR = Path(__file__).resolve().parents[1]
MOM_SAMPLE = Path(
    r"F:\Collage\Sem5\PBL\TemplateOS\backend\storage\templates\original"
    r"\t3-format-mom-updated-d38ce0f9440849c39db4b87b2c572b90.docx"
)
DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)


@event.listens_for(test_engine, "connect")
def _enable_foreign_keys(dbapi_connection, _record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


TestSession = sessionmaker(bind=test_engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)

OWNER_EMAIL = "gen_owner@example.com"
STRANGER_EMAIL = "gen_stranger@example.com"


@pytest.fixture(scope="module", autouse=True)
def setup_db(tmp_path_factory):
    old = app.dependency_overrides.get(get_db)
    old_root = storage_service.root_path
    storage_service.root_path = tmp_path_factory.mktemp("gen-storage").resolve()
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)
    storage_service.root_path = old_root
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


@pytest.fixture(autouse=True)
def pdf_off(monkeypatch):
    """No PDF converter unless a test configures one (never a real soffice)."""
    monkeypatch.setattr(settings, "pdf_converter", "none")
    monkeypatch.setattr(settings, "gotenberg_url", None)
    monkeypatch.setattr(settings, "libreoffice_path", None)


def _signup_and_login(email: str, full_name: str) -> str:
    client.post(
        "/api/v1/auth/signup",
        json={"full_name": full_name, "email": email, "password": "strong-password"},
    )
    login = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "strong-password"}
    )
    assert login.status_code == 200, login.text
    return login.json()["access_token"]


@pytest.fixture(scope="module")
def owner_token(setup_db):
    return _signup_and_login(OWNER_EMAIL, "Gen Owner")


@pytest.fixture(scope="module")
def stranger_token(setup_db):
    return _signup_and_login(STRANGER_EMAIL, "Gen Stranger")


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


# ── DOCX helpers ──────────────────────────────────────────────────


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _docx_text(docx_bytes: bytes) -> str:
    """All visible text of body, headers and footers (paragraph per line)."""
    lines = []
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
        for name in archive.namelist():
            if not re.match(r"word/(document|header\d*|footer\d*)\.xml$", name):
                continue
            xml = archive.read(name).decode("utf-8")
            for paragraph in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.DOTALL):
                paragraph = re.sub(r"<w:br/>", "\n", paragraph)
                texts = re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>|(\n)", paragraph)
                lines.append(html.unescape("".join(t or nl for t, nl in texts)))
    return "\n".join(lines)


SIMPLE_TEMPLATE = [
    "Title: {{ title }}",
    "Date: {{ meeting_date }}",
    "Count: {{ count }}",
    "Notes: {{ notes }}",
    "{%p for item in agenda_items %}",
    "- {{ item }}",
    "{%p endfor %}",
    "End.",
]


def _field(name: str, field_type: str, required: bool = False) -> TemplateField:
    return TemplateField(field_name=name, field_type=field_type, is_required=required)


SIMPLE_FIELDS = [
    ("title", "text", True),
    ("meeting_date", "date", False),
    ("count", "number", False),
    ("notes", "textarea", False),
    ("agenda_items", "list", True),
]


# ── 1. generation_service ────────────────────────────────────────


def test_format_date_value():
    assert format_date_value("2026-10-13") == "13-10-2026"
    assert format_date_value("13/10/2026") == "13/10/2026"
    assert format_date_value("2026-02-30") == "2026-02-30"
    assert format_date_value("") == ""


def test_build_context_types_and_empty_values():
    fields = [_field(name, ftype) for name, ftype, _ in SIMPLE_FIELDS]
    context = build_context(
        fields,
        {
            "title": "A & B <x>",
            "meeting_date": "2026-10-13",
            "count": "42",
            "agenda_items": json.dumps(["One", " ", "Two"]),
            "notes": None,
        },
    )
    assert context["title"] == "A & B <x>"
    assert context["meeting_date"] == "13-10-2026"
    assert context["count"] == "42"
    assert context["notes"] == ""
    assert isinstance(context["agenda_items"], list)
    assert context["agenda_items"] == ["One", "Two"]
    assert str(context["agenda_items"]) == "One\nTwo"


def test_build_context_list_edge_cases():
    fields = [_field("items", "list")]
    assert build_context(fields, {})["items"] == []
    assert build_context(fields, {"items": ""})["items"] == []
    assert build_context(fields, {"items": "not json"})["items"] == ["not json"]


def test_render_lists_dates_and_empty_optionals():
    fields = [_field(name, ftype) for name, ftype, _ in SIMPLE_FIELDS]
    context = build_context(
        fields,
        {
            "title": "Budget & Plans",
            "meeting_date": "2026-10-13",
            "count": "7",
            "agenda_items": json.dumps(["Welcome", "Budget review", "AOB"]),
        },
    )
    output = render_docx(_docx_bytes(SIMPLE_TEMPLATE), context)
    text = _docx_text(output)
    assert "Title: Budget & Plans" in text
    assert "Date: 13-10-2026" in text
    assert "Count: 7" in text
    assert "- Welcome\n- Budget review\n- AOB" in text
    assert "Notes: \n" in text or text.split("Notes:")[1].startswith("\n")
    assert "None" not in text
    assert "{{" not in text and "{%" not in text


def test_render_multiline_text_becomes_line_breaks():
    output = render_docx(
        _docx_bytes(["{{ notes }}"]),
        build_context([_field("notes", "textarea")], {"notes": "line 1\r\nline 2"}),
    )
    assert "<w:br/>" in zipfile.ZipFile(io.BytesIO(output)).read("word/document.xml").decode()


def test_render_undefined_variable_is_a_clear_error():
    with pytest.raises(TemplateRenderError) as info:
        render_docx(_docx_bytes(["Hello {{ missing_key }}"]), {"title": "x"})
    assert "missing_key" in str(info.value)


def test_render_syntax_error_is_a_clear_error():
    with pytest.raises(TemplateRenderError) as info:
        render_docx(_docx_bytes(["{%p for item in %}", "{%p endfor %}"]), {})
    assert "syntax error" in str(info.value)


def test_download_filename_is_sanitized():
    doc = Document(id=7, name='Weekly MoM: 13/10 <final>?')
    assert download_filename(doc, "docx") == "Weekly MoM 13 10 final.docx"
    assert download_filename(Document(id=7, name="///"), "pdf") == "document-7.pdf"
    assert download_filename(Document(id=7, name=None), "pdf") == "document-7.pdf"


# ── 2. Endpoint helpers ───────────────────────────────────────────


def _owner_id() -> int:
    with TestSession() as session:
        return session.execute(
            select(User.id).where(User.email == OWNER_EMAIL)
        ).scalar_one()


def _make_template(paragraphs: list[str], fields=SIMPLE_FIELDS, name="Gen Template") -> int:
    stored = storage_service.save_bytes(
        "templates_processed", "gen.docx", _docx_bytes(paragraphs)
    )
    with TestSession() as session:
        template = Template(
            name=name,
            category="mom",
            visibility="public",
            uploaded_by=_owner_id(),
            status="field_configured",
            processed_file_path=stored.path,
        )
        session.add(template)
        session.commit()
        template_id = template.id
        for field_name, field_type, required in fields:
            create_field(
                session,
                template_id,
                TemplateFieldCreate(
                    field_name=field_name, field_type=field_type, is_required=required
                ),
            )
    return template_id


def _create_doc(token: str, template_id: int, name: str = "Weekly MoM: 13/10") -> int:
    response = client.post(
        "/api/v1/documents/create",
        headers=_auth(token),
        json={"template_id": template_id, "name": name},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _save(token: str, document_id: int, values: dict[str, str]) -> None:
    response = client.put(
        f"/api/v1/documents/{document_id}/values",
        headers=_auth(token),
        json={"values": [{"field_name": k, "value": v} for k, v in values.items()]},
    )
    assert response.status_code == 200, response.text


GOOD_VALUES = {
    "title": "Budget review",
    "meeting_date": "2026-10-13",
    "agenda_items": json.dumps(["Welcome", "Budget"]),
}


def _generate(token: str, document_id: int):
    return client.post(f"/api/v1/documents/{document_id}/generate", headers=_auth(token))


@pytest.fixture()
def ready_doc(owner_token) -> int:
    template_id = _make_template(SIMPLE_TEMPLATE)
    document_id = _create_doc(owner_token, template_id)
    _save(owner_token, document_id, GOOD_VALUES)
    return document_id


def _record(document_id: int) -> GeneratedFile | None:
    with TestSession() as session:
        return session.execute(
            select(GeneratedFile).where(GeneratedFile.document_id == document_id)
        ).scalar_one_or_none()


# ── 3. POST /generate ─────────────────────────────────────────────


def test_generate_renders_and_records(owner_token, ready_doc):
    response = _generate(owner_token, ready_doc)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["document_id"] == ready_doc
    assert body["document_status"] == "generated"
    assert body["pdf_status"] == "not_configured"
    assert body["pdf_error"] is None
    assert body["pdf_url"] is None
    assert body["docx_url"] == f"/api/v1/documents/{ready_doc}/download?format=docx"
    assert body["generated_at"]
    assert body["is_stale"] is False

    record = _record(ready_doc)
    assert record is not None
    assert record.docx_path.startswith("generated/docx/")
    assert record.template_version == 1
    assert record.template_source_path.startswith("templates/processed/")
    text = _docx_text(storage_service.read_bytes(record.docx_path))
    assert "Title: Budget review" in text
    assert "Date: 13-10-2026" in text
    assert "- Welcome\n- Budget" in text
    assert "None" not in text

    detail = client.get(f"/api/v1/documents/{ready_doc}", headers=_auth(owner_token)).json()
    assert detail["status"] == "generated"
    assert detail["generated"]["pdf_status"] == "not_configured"
    assert detail["generated"]["is_stale"] is False

    listed = client.get("/api/v1/documents", headers=_auth(owner_token)).json()
    item = next(d for d in listed if d["id"] == ready_doc)
    assert item["generated"]["docx_url"].endswith("format=docx")


def test_list_and_detail_have_null_generated_before_generation(owner_token):
    template_id = _make_template(SIMPLE_TEMPLATE)
    document_id = _create_doc(owner_token, template_id)
    detail = client.get(f"/api/v1/documents/{document_id}", headers=_auth(owner_token))
    assert detail.json()["generated"] is None
    listed = client.get("/api/v1/documents", headers=_auth(owner_token)).json()
    assert next(d for d in listed if d["id"] == document_id)["generated"] is None


def test_saving_after_generation_marks_stale(owner_token, ready_doc):
    assert _generate(owner_token, ready_doc).status_code == 200
    _save(owner_token, ready_doc, {"title": "Changed title"})
    detail = client.get(f"/api/v1/documents/{ready_doc}", headers=_auth(owner_token)).json()
    assert detail["generated"]["is_stale"] is True
    # Regenerating clears it.
    body = _generate(owner_token, ready_doc).json()
    assert body["is_stale"] is False


def test_generate_missing_required_values_422(owner_token):
    template_id = _make_template(SIMPLE_TEMPLATE)
    document_id = _create_doc(owner_token, template_id)
    _save(owner_token, document_id, {"meeting_date": "2026-10-13"})
    response = _generate(owner_token, document_id)
    assert response.status_code == 422
    errors = {e["field_name"]: e["error"] for e in response.json()["detail"]}
    assert errors == {
        "title": "This field is required",
        "agenda_items": "This field is required",
    }
    assert _record(document_id) is None
    detail = client.get(f"/api/v1/documents/{document_id}", headers=_auth(owner_token))
    assert detail.json()["status"] == "draft"


def test_generate_template_error_422(owner_token):
    template_id = _make_template(["Hi {{ title }} {{ not_a_field }}"], fields=[("title", "text", True)])
    document_id = _create_doc(owner_token, template_id)
    _save(owner_token, document_id, {"title": "x"})
    response = _generate(owner_token, document_id)
    assert response.status_code == 422
    assert isinstance(response.json()["detail"], str)
    assert "not_a_field" in response.json()["detail"]
    assert _record(document_id) is None


def test_generate_template_file_missing_409(owner_token):
    template_id = _make_template(SIMPLE_TEMPLATE)
    document_id = _create_doc(owner_token, template_id)
    _save(owner_token, document_id, GOOD_VALUES)
    with TestSession() as session:
        storage_service.delete_file(session.get(Template, template_id).processed_file_path)
    assert _generate(owner_token, document_id).status_code == 409


def test_generation_is_owner_only(owner_token, stranger_token, ready_doc):
    assert _generate(stranger_token, ready_doc).status_code == 403
    assert _generate(owner_token, ready_doc).status_code == 200
    for path in ("generate-pdf",):
        assert (
            client.post(f"/api/v1/documents/{ready_doc}/{path}", headers=_auth(stranger_token)).status_code
            == 403
        )
    download = client.get(
        f"/api/v1/documents/{ready_doc}/download", headers=_auth(stranger_token)
    )
    assert download.status_code == 403
    assert _generate(owner_token, 999_999).status_code == 404
    assert client.post(f"/api/v1/documents/{ready_doc}/generate").status_code == 401


# ── 4. Download ───────────────────────────────────────────────────


def test_download_docx_headers_and_filename(owner_token, ready_doc):
    assert _generate(owner_token, ready_doc).status_code == 200
    response = client.get(
        f"/api/v1/documents/{ready_doc}/download?format=docx", headers=_auth(owner_token)
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == DOCX_TYPE
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment")
    assert 'filename="Weekly MoM 13 10.docx"' in disposition
    assert response.headers["cache-control"] == "no-store"
    assert "Title: Budget review" in _docx_text(response.content)


def test_download_not_generated_404_and_pdf_unavailable_409(owner_token, ready_doc):
    url = f"/api/v1/documents/{ready_doc}/download"
    assert client.get(url, headers=_auth(owner_token)).status_code == 404
    assert _generate(owner_token, ready_doc).status_code == 200
    response = client.get(f"{url}?format=pdf", headers=_auth(owner_token))
    assert response.status_code == 409
    assert "PDF" in response.json()["detail"]
    assert client.get(f"{url}?format=xlsx", headers=_auth(owner_token)).status_code == 422


def test_download_missing_file_409(owner_token, ready_doc):
    assert _generate(owner_token, ready_doc).status_code == 200
    storage_service.delete_file(_record(ready_doc).docx_path)
    response = client.get(
        f"/api/v1/documents/{ready_doc}/download", headers=_auth(owner_token)
    )
    assert response.status_code == 409


# ── 5. PDF conversion ─────────────────────────────────────────────

FAKE_PDF = b"%PDF-1.7\n% fake pdf for tests\n%%EOF"


@pytest.fixture()
def gotenberg(monkeypatch):
    """Configure Gotenberg and capture requests; set .response / .error."""

    class FakeGotenberg:
        response = httpx.Response(200, content=FAKE_PDF)
        error: Exception | None = None
        calls: list[dict] = []

    fake = FakeGotenberg()
    fake.calls = []

    def fake_post(url, files=None, timeout=None):
        fake.calls.append({"url": url, "files": files, "timeout": timeout})
        if fake.error is not None:
            raise fake.error
        return fake.response

    monkeypatch.setattr(settings, "pdf_converter", "auto")
    monkeypatch.setattr(settings, "gotenberg_url", "http://gotenberg.test:3000/")
    monkeypatch.setattr(settings, "gotenberg_timeout_seconds", 12)
    monkeypatch.setattr(pdf_service.httpx, "post", fake_post)
    return fake


def test_generate_with_gotenberg_success(owner_token, ready_doc, gotenberg):
    body = _generate(owner_token, ready_doc).json()
    assert body["pdf_status"] == "ok"
    assert body["pdf_url"] == f"/api/v1/documents/{ready_doc}/download?format=pdf"
    call = gotenberg.calls[0]
    assert call["url"] == "http://gotenberg.test:3000/forms/libreoffice/convert"
    assert call["timeout"] == 12
    filename, content, media_type = call["files"]["files"]
    assert filename.endswith(".docx") and media_type == DOCX_TYPE
    assert content.startswith(b"PK")

    record = _record(ready_doc)
    assert record.pdf_path.startswith("generated/pdf/")
    response = client.get(
        f"/api/v1/documents/{ready_doc}/download?format=pdf", headers=_auth(owner_token)
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert 'filename="Weekly MoM 13 10.pdf"' in response.headers["content-disposition"]
    assert response.content == FAKE_PDF


def test_gotenberg_http_failure_keeps_docx(owner_token, ready_doc, gotenberg):
    gotenberg.response = httpx.Response(503, text="busy")
    response = _generate(owner_token, ready_doc)
    assert response.status_code == 200
    body = response.json()
    assert body["document_status"] == "generated"
    assert body["pdf_status"] == "failed"
    assert "503" in body["pdf_error"]
    assert body["pdf_url"] is None
    record = _record(ready_doc)
    assert record.pdf_path is None and record.pdf_error
    assert storage_service.exists(record.docx_path)


def test_gotenberg_timeout_then_retry_succeeds(owner_token, ready_doc, gotenberg):
    gotenberg.error = httpx.ReadTimeout("slow")
    body = _generate(owner_token, ready_doc).json()
    assert body["pdf_status"] == "failed"
    assert "timed out" in body["pdf_error"]

    gotenberg.error = None
    retry = client.post(
        f"/api/v1/documents/{ready_doc}/generate-pdf", headers=_auth(owner_token)
    )
    assert retry.status_code == 200, retry.text
    assert retry.json()["pdf_status"] == "ok"
    assert retry.json()["pdf_error"] is None
    record = _record(ready_doc)
    assert record.pdf_error is None and storage_service.exists(record.pdf_path)


def test_gotenberg_non_pdf_response_is_a_failure(owner_token, ready_doc, gotenberg):
    gotenberg.response = httpx.Response(200, content=b"<html>oops</html>")
    assert _generate(owner_token, ready_doc).json()["pdf_status"] == "failed"


def test_retry_failure_and_not_configured(owner_token, ready_doc, gotenberg, monkeypatch):
    gotenberg.response = httpx.Response(500)
    assert _generate(owner_token, ready_doc).json()["pdf_status"] == "failed"
    retry = client.post(
        f"/api/v1/documents/{ready_doc}/generate-pdf", headers=_auth(owner_token)
    )
    assert retry.status_code == 200
    assert retry.json()["pdf_status"] == "failed"

    monkeypatch.setattr(settings, "pdf_converter", "none")
    retry = client.post(
        f"/api/v1/documents/{ready_doc}/generate-pdf", headers=_auth(owner_token)
    )
    assert retry.json()["pdf_status"] == "not_configured"
    assert retry.json()["pdf_error"] is None


def test_retry_before_generation_404(owner_token, ready_doc):
    response = client.post(
        f"/api/v1/documents/{ready_doc}/generate-pdf", headers=_auth(owner_token)
    )
    assert response.status_code == 404


def test_converter_selection(monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "pdf_converter", "auto")
    monkeypatch.setattr(pdf_service.shutil, "which", lambda _name: None)
    assert pdf_service.get_pdf_converter() is None

    monkeypatch.setattr(settings, "gotenberg_url", "http://g:3000")
    assert pdf_service.get_pdf_converter().name == "gotenberg"

    binary = tmp_path / "soffice.exe"
    binary.write_bytes(b"")
    monkeypatch.setattr(settings, "libreoffice_path", str(binary))
    monkeypatch.setattr(settings, "pdf_converter", "libreoffice")
    converter = pdf_service.get_pdf_converter()
    assert converter.name == "libreoffice" and converter.binary == str(binary)

    monkeypatch.setattr(settings, "gotenberg_url", None)
    monkeypatch.setattr(settings, "pdf_converter", "auto")
    assert pdf_service.get_pdf_converter().name == "libreoffice"

    monkeypatch.setattr(settings, "libreoffice_path", str(tmp_path / "missing.exe"))
    assert pdf_service.get_pdf_converter() is None

    monkeypatch.setattr(settings, "gotenberg_url", "http://g:3000")
    monkeypatch.setattr(settings, "pdf_converter", "none")
    assert pdf_service.get_pdf_converter() is None


def test_libreoffice_converter(monkeypatch):
    seen = {}

    def fake_run(command, capture_output, timeout, check):
        seen["command"] = command
        seen["timeout"] = timeout
        outdir = Path(command[command.index("--outdir") + 1])
        assert Path(command[-1]).read_bytes() == b"PK-docx"
        (outdir / "document.pdf").write_bytes(FAKE_PDF)
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(pdf_service.subprocess, "run", fake_run)
    converter = pdf_service.LibreOfficeConverter("soffice", 30)
    assert converter.convert(b"PK-docx", "x.docx") == FAKE_PDF
    assert "--headless" in seen["command"] and "pdf" in seen["command"]
    assert seen["timeout"] == 30

    def failing_run(command, capture_output, timeout, check):
        return subprocess.CompletedProcess(command, 1, b"", b"boom")

    monkeypatch.setattr(pdf_service.subprocess, "run", failing_run)
    with pytest.raises(pdf_service.PdfConversionError):
        converter.convert(b"PK-docx", "x.docx")

    def slow_run(command, capture_output, timeout, check):
        raise subprocess.TimeoutExpired(command, timeout)

    monkeypatch.setattr(pdf_service.subprocess, "run", slow_run)
    with pytest.raises(pdf_service.PdfConversionError, match="timed out"):
        converter.convert(b"PK-docx", "x.docx")


# ── 6. Regeneration and cleanup ───────────────────────────────────


def test_regeneration_replaces_files(owner_token, ready_doc, gotenberg):
    assert _generate(owner_token, ready_doc).status_code == 200
    first = _record(ready_doc)
    assert _generate(owner_token, ready_doc).status_code == 200
    second = _record(ready_doc)
    assert second.id == first.id  # upsert, still one row
    assert second.docx_path != first.docx_path
    assert second.pdf_path != first.pdf_path
    assert not storage_service.exists(first.docx_path)
    assert not storage_service.exists(first.pdf_path)
    assert storage_service.exists(second.docx_path)
    assert storage_service.exists(second.pdf_path)


def test_document_delete_cascades_row_and_files(owner_token, ready_doc, gotenberg):
    assert _generate(owner_token, ready_doc).status_code == 200
    record = _record(ready_doc)
    with TestSession() as session:
        assert delete_document(session, ready_doc) is True
    assert _record(ready_doc) is None
    assert not storage_service.exists(record.docx_path)
    assert not storage_service.exists(record.pdf_path)


def test_template_delete_cascades_generated_files(owner_token, ready_doc):
    assert _generate(owner_token, ready_doc).status_code == 200
    record = _record(ready_doc)
    with TestSession() as session:
        template_id = session.get(Document, ready_doc).template_id
        session.delete(session.get(Template, template_id))
        session.commit()
    assert _record(ready_doc) is None
    assert not storage_service.exists(record.docx_path)


def test_failed_delete_keeps_files(owner_token, ready_doc):
    assert _generate(owner_token, ready_doc).status_code == 200
    record = _record(ready_doc)
    with TestSession() as session:
        session.delete(session.get(Document, ready_doc))
        session.flush()
        session.rollback()
    assert _record(ready_doc) is not None
    assert storage_service.exists(record.docx_path)


# ── 7. Migration ──────────────────────────────────────────────────


def _alembic(environment: dict, *args: str) -> None:
    subprocess.run(
        [sys.executable, "-m", "alembic", "-c", "alembic.ini.example", *args],
        cwd=BACKEND_DIR,
        env=environment,
        check=True,
        capture_output=True,
        text=True,
    )


def test_generated_files_migration_up_down_up(tmp_path: Path) -> None:
    environment = {
        **os.environ,
        "DATABASE_URL": f"sqlite:///{(tmp_path / 'generated.db').as_posix()}",
        "JWT_SECRET_KEY": "test-secret-that-is-long-enough-for-migration-tests",
        "PYTHONDONTWRITEBYTECODE": "1",
    }

    def tables() -> set[str]:
        engine = create_engine(environment["DATABASE_URL"])
        try:
            return set(inspect(engine).get_table_names())
        finally:
            engine.dispose()

    _alembic(environment, "upgrade", "f1b7c3d9e2a4")
    assert "generated_files" not in tables()

    _alembic(environment, "upgrade", "b6d2e8f4a1c7")
    engine = create_engine(environment["DATABASE_URL"])
    try:
        inspector = inspect(engine)
        columns = {c["name"]: c for c in inspector.get_columns("generated_files")}
        assert set(columns) == {
            "id",
            "document_id",
            "docx_path",
            "pdf_path",
            "pdf_error",
            "template_version",
            "template_source_path",
            "generated_at",
        }
        assert columns["docx_path"]["nullable"] is False
        assert columns["pdf_path"]["nullable"] is True
        assert columns["pdf_error"]["nullable"] is True
        unique = {
            tuple(c["column_names"]) for c in inspector.get_unique_constraints("generated_files")
        }
        assert ("document_id",) in unique
        foreign = inspector.get_foreign_keys("generated_files")
        assert foreign[0]["referred_table"] == "documents"
        assert foreign[0]["options"].get("ondelete") == "CASCADE"
    finally:
        engine.dispose()

    _alembic(environment, "downgrade", "f1b7c3d9e2a4")
    assert "generated_files" not in tables()

    _alembic(environment, "upgrade", "b6d2e8f4a1c7")
    assert "generated_files" in tables()


# ── 8. End to end: real Minutes-of-Meeting sample ────────────────


def _mom_fixture_docx() -> bytes:
    """Fallback when the real sample is unavailable: a MoM-like document."""
    document = DocxDocument()
    document.add_paragraph("Ref No: SITN/2601/Nomenclature/Sr.No")
    document.add_paragraph("Date: _____________")
    document.add_paragraph("MINUTES OF MEETING")
    document.add_paragraph(
        "Meeting of the ------------- committee was held on --------- at --------- "
        "at ------ chaired by -------------."
    )
    document.add_paragraph("Agenda:")
    document.add_paragraph("ABC")
    document.add_paragraph("ABC")
    document.add_paragraph("Discussion:")
    document.add_paragraph("1.")
    document.add_paragraph("2.")
    table = document.add_table(rows=2, cols=3)
    for cell, label in zip(table.rows[0].cells, ("Prepared by", "Verified by", "Approved by")):
        cell.text = label
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _value_for(field: dict) -> str | None:
    name, field_type = field["field_name"], field["field_type"]
    rule = field.get("validation_rule") or ""
    if field_type == "list":
        return json.dumps([f"{name} item one", f"{name} item two & more"])
    if field_type == "date":
        return "2026-10-13"
    if field_type == "number":
        return "5"
    if "email" in rule:
        return "someone@example.com"
    value = f"Val {name}"
    match = re.search(r"max:(\d+)", rule)
    if match:
        value = value[: int(match.group(1))]
    return value


def test_end_to_end_real_mom_sample(owner_token):
    if MOM_SAMPLE.is_file():
        sample_bytes, used_real = MOM_SAMPLE.read_bytes(), True
    else:
        sample_bytes, used_real = _mom_fixture_docx(), False

    upload = client.post(
        "/api/v1/templates/upload",
        files={"file": ("mom.docx", io.BytesIO(sample_bytes), DOCX_TYPE)},
        data={"name": "MoM E2E", "category": "mom", "visibility": "private"},
        headers=_auth(owner_token),
    )
    assert upload.status_code == 201, upload.text
    template_id = upload.json()["id"]

    suggestions = client.get(
        f"/api/v1/templates/{template_id}/clean-suggestions", headers=_auth(owner_token)
    )
    assert suggestions.status_code == 200, suggestions.text
    items = suggestions.json()["suggestions"]
    assert items, "the cleaner found nothing to turn into placeholders"
    replacements = [
        {
            "sample_text": s["sample_text"],
            "anchor": s["anchor"],
            "placeholder_key": s["placeholder_key"],
            "field_label": s["field_label"],
            "field_type": s["field_type"],
        }
        for s in items
    ]
    cleaned = client.post(
        f"/api/v1/templates/{template_id}/clean",
        headers=_auth(owner_token),
        json={"replacements": replacements, "confirm": True, "mark_configured": True},
    )
    assert cleaned.status_code == 200, cleaned.text
    assert cleaned.json()["warnings"]["render_error"] is None

    fields = client.get(
        f"/api/v1/templates/{template_id}/fields", headers=_auth(owner_token)
    ).json()
    by_type = {f["field_name"]: f["field_type"] for f in fields}
    assert "list" in by_type.values()
    if used_real:
        assert by_type.get("agenda_items") == "list"

    # Fill every required field; leave the optional ones empty.
    values = {f["field_name"]: _value_for(f) for f in fields if f["is_required"]}
    optional = [f["field_name"] for f in fields if not f["is_required"]]

    document_id = _create_doc(owner_token, template_id, name="MoM 13 Oct")
    _save(owner_token, document_id, values)
    response = _generate(owner_token, document_id)
    assert response.status_code == 200, response.text
    assert response.json()["document_status"] == "generated"

    download = client.get(
        f"/api/v1/documents/{document_id}/download?format=docx", headers=_auth(owner_token)
    )
    assert download.status_code == 200
    assert 'filename="MoM 13 Oct.docx"' in download.headers["content-disposition"]
    DocxDocument(io.BytesIO(download.content))  # opens as a valid DOCX
    text = _docx_text(download.content)

    for name, value in values.items():
        if by_type[name] == "list":
            for item in json.loads(value):
                assert item in text, f"list item of {name} missing"
        elif by_type[name] == "date":
            assert "13-10-2026" in text
        else:
            assert value in text, f"{name} missing from output"
    assert "{{" not in text and "{%" not in text
    assert "None" not in text
    # Each list item rendered on its own paragraph (the {%p for %} loop).
    list_names = [n for n, t in by_type.items() if t == "list" and n in values]
    for name in list_names:
        first, second = json.loads(values[name])
        assert re.search(re.escape(first) + r".*\n.*" + re.escape(second), text)
    print(
        f"\nMoM e2e: real_sample={used_real}, fields={len(fields)}, "
        f"filled={len(values)}, optional_empty={optional}"
    )
