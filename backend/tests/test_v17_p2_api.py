"""
V1.7 Phase 2 API (Member 2): rich values, AI flags, rich rendering,
versioned generation, version endpoints.

Coverage:
- PUT /values: rich content normalized + plain value derived, required /
  min / max on the plain text, 422 for unsupported formatting, unsafe
  links, rich on a non-textarea field, missing content, ai_modified without
  an action; plain-only clients unchanged; GET returns the new fields.
- Generation: the DOCX carries bold runs, a real hyperlink, bullet / number
  lines, highlight, escaping, emoji; valid XML (no runs inside <w:t>); the
  placeholder run's font is kept; filters get plain text; invalid stored
  JSON falls back to plain text; plain textareas render as before.
- Versions: generate twice -> v1, v2 (newest first, changed_fields,
  created_by_name, is_latest, URLs); detail from the snapshot; per-version
  downloads (filename "<name> v<n>.<ext>", 404 / 409 / 410); retention
  (keep=1) prunes old files -> 410 while the detail stays 200; restore
  (values = snapshot, absent fields cleared, then v3), 409 when not
  editable; 403 for other users, 404 unknown version / document.
- Preview: per-field rich / AI fields and generation version counters.
- Document delete removes every version's files.

Storage is a temporary folder; the database is in-memory SQLite with
foreign keys on; PDF conversion is a fake converter (never a real service).
"""

import io
import json
import os
import re
import zipfile

import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from lxml import etree
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests")

from app.core.config import settings  # noqa: E402
from app.crud.document_crud import delete_document  # noqa: E402
from app.crud.template_field_crud import create_field  # noqa: E402
from app.db.base import Base  # noqa: E402
from app.db.session import get_db  # noqa: E402
from app.main import app  # noqa: E402
from app.models.document import Document  # noqa: E402
from app.models.document_value import DocumentValue  # noqa: E402
from app.models.generated_file import GeneratedFile  # noqa: E402
from app.models.template import Template  # noqa: E402
from app.models.user import User  # noqa: E402
from app.schemas.template_field import TemplateFieldCreate  # noqa: E402
from app.services import document_generation  # noqa: E402
from app.services.rich_text import to_plain_text, validate_rich_content  # noqa: E402
from app.services.storage_service import storage_service  # noqa: E402
from test_rich_text import SHARED_DOC, SHARED_PLAIN  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)


@event.listens_for(engine, "connect")
def _enable_foreign_keys(dbapi_connection, _record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)
OWNER = "p2_owner@example.com"
STRANGER = "p2_stranger@example.com"
FAKE_PDF = b"%PDF-1.4\n% fake pdf for tests\n" + b"0" * 600
DOCS = "/api/v1/documents"


@pytest.fixture(scope="module", autouse=True)
def setup_db(tmp_path_factory):
    old = app.dependency_overrides.get(get_db)
    old_root = storage_service.root_path
    storage_service.root_path = tmp_path_factory.mktemp("p2-storage").resolve()
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    storage_service.root_path = old_root
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


class FakeConverter:
    def __init__(self):
        self.calls = 0

    def convert(self, docx_bytes: bytes, filename: str) -> bytes:
        self.calls += 1
        return FAKE_PDF


@pytest.fixture(autouse=True)
def fake_pdf(monkeypatch):
    converter = FakeConverter()
    monkeypatch.setattr(document_generation, "get_pdf_converter", lambda: converter)
    monkeypatch.setattr(settings, "generated_versions_keep_files", 10)
    return converter


def _login(email: str, name: str) -> str:
    client.post("/api/v1/auth/signup", json={"full_name": name, "email": email, "password": "strong-password"})
    response = client.post("/api/v1/auth/login", json={"email": email, "password": "strong-password"})
    assert response.status_code == 200, response.text
    return response.json()["access_token"]


@pytest.fixture(scope="module")
def owner(setup_db) -> str:
    return _login(OWNER, "Priya Owner")


@pytest.fixture(scope="module")
def stranger(setup_db) -> str:
    return _login(STRANGER, "Sam Stranger")


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def _docx_template(paragraphs: list) -> bytes:
    """paragraphs: strings, or lists of (text, {font/bold}) runs."""
    document = DocxDocument()
    for paragraph in paragraphs:
        if isinstance(paragraph, str):
            document.add_paragraph(paragraph)
            continue
        p = document.add_paragraph()
        for text, style in paragraph:
            run = p.add_run(text)
            run.bold = style.get("bold")
            if style.get("font"):
                run.font.name = style["font"]
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


TEMPLATE = [
    "Title: {{ title }}",
    [("Discussion: ", {}), ("{{ discussion }}", {"font": "Arial"}), (" (end)", {})],
    "Shout: {{ discussion|upper }}",
    "{% if notes %}Notes: {{ notes }}{% endif %}",
]
FIELDS = [
    ("title", "text", True, None),
    ("discussion", "textarea", True, "max:400"),
    ("notes", "textarea", False, None),
    ("short", "textarea", False, "min:3|max:12"),
]


def _owner_id() -> int:
    with TestSession() as session:
        return session.execute(select(User.id).where(User.email == OWNER)).scalar_one()


def _make_template(paragraphs=TEMPLATE, fields=FIELDS) -> int:
    stored = storage_service.save_bytes("templates_processed", "p2.docx", _docx_template(paragraphs))
    with TestSession() as session:
        template = Template(
            name="P2 Template",
            category="mom",
            visibility="public",
            uploaded_by=_owner_id(),
            status="field_configured",
            processed_file_path=stored.path,
        )
        session.add(template)
        session.commit()
        for name, ftype, required, rule in fields:
            create_field(
                session,
                template.id,
                TemplateFieldCreate(field_name=name, field_type=ftype, is_required=required, validation_rule=rule),
            )
        return template.id


def _create(token: str, template_id: int, name: str = "Weekly MoM") -> int:
    response = client.post(f"{DOCS}/create", headers=_auth(token), json={"template_id": template_id, "name": name})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _put(token: str, document_id: int, values: list[dict]):
    return client.put(f"{DOCS}/{document_id}/values", headers=_auth(token), json={"values": values})


def _rich(name: str, doc: dict, **extra) -> dict:
    return {"field_name": name, "value": "ignored by the server", "value_format": "rich", "rich_content": doc, **extra}


def _text(value: str) -> dict:
    return {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": value}]}]}


@pytest.fixture()
def draft(owner) -> int:
    document_id = _create(owner, _make_template())
    response = _put(
        owner,
        document_id,
        [
            {"field_name": "title", "value": "Budget"},
            _rich("discussion", SHARED_DOC, ai_modified=True, ai_action="formal"),
        ],
    )
    assert response.status_code == 200, response.text
    return document_id


def _generate(token: str, document_id: int):
    return client.post(f"{DOCS}/{document_id}/generate", headers=_auth(token))


def _versions(token: str, document_id: int):
    return client.get(f"{DOCS}/{document_id}/versions", headers=_auth(token))


def _download(token: str, document_id: int, version: int, fmt: str):
    return client.get(f"{DOCS}/{document_id}/versions/{version}/download/{fmt}", headers=_auth(token))


def _document_xml(docx_bytes: bytes) -> str:
    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode("utf-8")


def _paragraph_texts(docx_bytes: bytes) -> list[str]:
    return [p.text for p in DocxDocument(io.BytesIO(docx_bytes)).paragraphs]


# ── Save values ───────────────────────────────────────────────────


def test_save_rich_value_normalizes_and_derives_plain(owner, draft):
    response = client.get(f"{DOCS}/{draft}/values", headers=_auth(owner))
    assert response.status_code == 200
    values = {v["field_name"]: v for v in response.json()["values"]}
    discussion = values["discussion"]
    assert discussion["value"] == SHARED_PLAIN
    assert discussion["value_format"] == "rich"
    assert discussion["rich_content"] == validate_rich_content(SHARED_DOC)
    assert "target" not in json.dumps(discussion["rich_content"])  # attrs dropped
    assert (discussion["ai_modified"], discussion["ai_action"]) == (True, "formal")
    assert discussion["ai_modified_at"] is not None
    title = values["title"]
    assert (title["value_format"], title["rich_content"], title["ai_modified"], title["ai_action"]) == (
        "plain", None, False, None,
    )
    detail = client.get(f"{DOCS}/{draft}", headers=_auth(owner)).json()
    assert {v["field_name"]: v["value_format"] for v in detail["values"]} == {"title": "plain", "discussion": "rich"}


def test_manual_edit_clears_ai_flag(owner, draft):
    response = _put(owner, draft, [_rich("discussion", _text("Typed by hand"))])
    assert response.status_code == 200, response.text
    saved = response.json()["values"][0]
    assert (saved["ai_modified"], saved["ai_action"], saved["ai_modified_at"]) == (False, None, None)
    assert saved["value"] == "Typed by hand"


@pytest.mark.parametrize(
    "item, error",
    [
        (_rich("discussion", {"type": "doc", "content": [{"type": "heading", "content": [{"type": "text", "text": "x"}]}]}),
         "Unsupported formatting: heading"),
        (_rich("discussion", {"type": "doc", "content": [{"type": "paragraph", "content": [
            {"type": "text", "text": "x", "marks": [{"type": "link", "attrs": {"href": "javascript:alert(1)"}}]}]}]}),
         "Links must start with http://, https:// or mailto:"),
        (_rich("title", _text("rich title")), "Rich text is only available for long text fields"),
        ({"field_name": "discussion", "value": "x", "value_format": "rich"}, "Rich text content is missing"),
        ({"field_name": "discussion", "value": "x", "ai_modified": True}, "ai_action is required when ai_modified is true"),
        (_rich("discussion", {"type": "doc", "content": []}), "This field is required"),
        (_rich("short", _text("ab")), "Must be at least 3 characters"),
        (_rich("short", {"type": "doc", "content": [{"type": "bulletList", "content": [
            {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "abcdefghij"}]}]}]}]}),
         "Must be at most 12 characters"),  # "• abcdefghij" = 12 chars is fine...
    ],
)
def test_save_rejects_invalid_rich_values(owner, draft, item, error):
    if error == "Must be at most 12 characters":
        # ...one more character in the bullet goes over: the bullet marker counts.
        item["rich_content"]["content"][0]["content"][0]["content"][0]["content"][0]["text"] += "k"
    response = _put(owner, draft, [item])
    assert response.status_code == 422, response.text
    assert {"field_name": item["field_name"], "error": error} in response.json()["detail"]


def test_save_rejects_unknown_ai_action_and_format(owner, draft):
    assert _put(owner, draft, [{"field_name": "discussion", "value": "x", "ai_modified": True, "ai_action": "rewrite"}]).status_code == 422
    assert _put(owner, draft, [{"field_name": "discussion", "value": "x", "value_format": "html"}]).status_code == 422


def test_plain_only_clients_unchanged(owner, draft):
    response = _put(owner, draft, [{"field_name": "discussion", "value": "Plain paragraph\nsecond line"}])
    assert response.status_code == 200
    saved = response.json()["values"][0]
    assert saved["value"] == "Plain paragraph\nsecond line"
    assert (saved["value_format"], saved["rich_content"]) == ("plain", None)
    body = _generate(owner, draft).json()
    record = _latest(draft)
    docx = storage_service.read_bytes(record.docx_path)
    assert "Discussion: Plain paragraph\nsecond line (end)" in _paragraph_texts(docx)
    assert body["version"] == 1


# ── Rich rendering ────────────────────────────────────────────────


def _latest(document_id: int) -> GeneratedFile:
    from app.crud.generated_file_crud import latest_generated_file

    with TestSession() as session:
        return latest_generated_file(session, document_id)


def test_generate_renders_rich_formatting(owner, draft, fake_pdf):
    response = _generate(owner, draft)
    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["version"], body["versions_count"], body["pdf_status"]) == (1, 1, "converted")
    docx = storage_service.read_bytes(_latest(draft).docx_path)
    xml = _document_xml(docx)
    root = etree.fromstring(xml.encode("utf-8"))  # well-formed
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    assert all(len(t) == 0 for t in root.iter(f"{w}t")), "elements nested in <w:t>"
    assert not root.xpath("//w:r[.//w:r]", namespaces={"w": w[1:-1]}), "runs nested in runs"
    assert "<w:b/>" in xml and "<w:i/>" in xml
    assert '<w:highlight w:val="yellow"/>' in xml
    assert "<w:hyperlink " in xml
    rels = zipfile.ZipFile(io.BytesIO(docx)).read("word/_rels/document.xml.rels").decode()
    assert 'Target="https://example.com" TargetMode="External"' in rels
    # The placeholder's own run font is kept for every rich run.
    discussion_runs = re.findall(r"<w:r>(?:(?!</w:r>).)*Apples(?:(?!</w:r>).)*</w:r>", xml, flags=re.DOTALL)
    assert discussion_runs and 'w:ascii="Arial"' in discussion_runs[0]
    texts = _paragraph_texts(docx)
    assert f"Discussion: {SHARED_PLAIN} (end)" in texts
    assert "Fourth &amp; &lt;last&gt;" in xml
    # Filters and tests see the plain text; empty optional rich -> falsy.
    assert f"Shout: {SHARED_PLAIN.upper()}" in texts
    assert not any(text.startswith("Notes:") for text in texts)
    assert fake_pdf.calls == 1


def test_rich_text_with_special_characters_and_emoji(owner):
    document_id = _create(owner, _make_template())
    doc = {"type": "doc", "content": [{"type": "paragraph", "content": [
        {"type": "text", "text": "R&D <beta> 🎉 \"q\"", "marks": [{"type": "bold"}, {"type": "underline"}, {"type": "strike"}]}]}]}
    assert _put(owner, document_id, [{"field_name": "title", "value": "T"}, _rich("discussion", doc), _rich("notes", _text("n"))]).status_code == 200
    assert _generate(owner, document_id).status_code == 200
    docx = storage_service.read_bytes(_latest(document_id).docx_path)
    xml = _document_xml(docx)
    etree.fromstring(xml.encode("utf-8"))
    assert "R&amp;D &lt;beta&gt; 🎉" in xml
    assert '<w:u w:val="single"/>' in xml and "<w:strike/>" in xml
    texts = _paragraph_texts(docx)
    assert 'Discussion: R&D <beta> 🎉 "q" (end)' in texts
    assert "Notes: n" in texts


def test_invalid_stored_rich_content_falls_back_to_plain(owner, draft):
    with TestSession() as session:
        row = session.execute(
            select(DocumentValue).where(DocumentValue.document_id == draft, DocumentValue.field_name == "discussion")
        ).scalar_one()
        row.rich_content = json.dumps({"type": "doc", "content": [{"type": "heading"}]})
        session.commit()
    assert _generate(owner, draft).status_code == 200
    texts = _paragraph_texts(storage_service.read_bytes(_latest(draft).docx_path))
    assert f"Discussion: {SHARED_PLAIN} (end)" in texts


# ── Versions ──────────────────────────────────────────────────────


def test_generate_twice_lists_versions_with_changes(owner, draft):
    assert _generate(owner, draft).json()["version"] == 1
    assert _put(owner, draft, [{"field_name": "title", "value": "Budget v2"}]).status_code == 200
    second = _generate(owner, draft).json()
    assert (second["version"], second["versions_count"]) == (2, 2)

    response = _versions(owner, draft)
    assert response.status_code == 200
    v2, v1 = response.json()
    assert (v2["version"], v2["is_latest"], v2["changed_fields"]) == (2, True, ["title"])
    assert (v1["version"], v1["is_latest"], v1["changed_fields"]) == (1, False, [])
    assert v2["created_by_name"] == "Priya Owner"
    assert (v2["has_docx"], v2["has_pdf"], v2["files_pruned"], v2["pdf_status"]) == (True, True, False, "converted")
    assert v1["docx_url"] == f"/api/v1/documents/{draft}/versions/1/download/docx"
    assert v1["pdf_url"] == f"/api/v1/documents/{draft}/versions/1/download/pdf"

    # Detail of v1 comes from its snapshot, not the current values.
    detail = client.get(f"{DOCS}/{draft}/versions/1", headers=_auth(owner)).json()
    fields = {f["field_name"]: f for s in detail["sections"] for f in s["fields"]}
    assert fields["title"]["display_value"] == "Budget"
    assert fields["discussion"]["value_format"] == "rich"
    assert fields["discussion"]["rich_content"] == validate_rich_content(SHARED_DOC)
    assert fields["discussion"]["display_value"] == SHARED_PLAIN
    assert (fields["discussion"]["ai_modified"], fields["discussion"]["ai_action"]) == (True, "formal")
    assert fields["notes"]["is_empty"] is True
    assert (detail["version"], detail["document_id"], detail["is_latest"]) == (1, draft, False)
    v2_detail = client.get(f"{DOCS}/{draft}/versions/2", headers=_auth(owner)).json()
    assert v2_detail["changed_fields"] == ["title"] and v2_detail["is_latest"] is True

    # Latest downloads and the document detail point at v2.
    assert client.get(f"{DOCS}/{draft}", headers=_auth(owner)).json()["generated"]["version"] == 2
    listed = client.get(DOCS, headers=_auth(owner)).json()
    assert next(d for d in listed if d["id"] == draft)["generated"]["version"] == 2


def test_version_downloads(owner, draft):
    _generate(owner, draft)
    _put(owner, draft, [{"field_name": "title", "value": "Second"}])
    _generate(owner, draft)
    v1 = _download(owner, draft, 1, "docx")
    assert v1.status_code == 200
    assert v1.headers["content-disposition"].startswith('attachment; filename="Weekly MoM v1.docx"')
    assert v1.headers["cache-control"] == "no-store"
    assert "Title: Budget" in _paragraph_texts(v1.content)
    assert "Title: Second" in _paragraph_texts(_download(owner, draft, 2, "docx").content)
    pdf = _download(owner, draft, 2, "pdf")
    assert pdf.status_code == 200 and pdf.content == FAKE_PDF
    assert pdf.headers["content-type"] == "application/pdf"
    assert 'filename="Weekly MoM v2.pdf"' in pdf.headers["content-disposition"]
    # Latest download is unchanged in shape and serves v2.
    latest = client.get(f"{DOCS}/{draft}/download?format=docx", headers=_auth(owner))
    assert "Title: Second" in _paragraph_texts(latest.content)
    assert _download(owner, draft, 9, "docx").status_code == 404
    assert _download(owner, draft, 1, "txt").status_code == 422


def test_version_download_without_pdf_and_missing_file(owner, draft, monkeypatch):
    monkeypatch.setattr(document_generation, "get_pdf_converter", lambda: None)
    assert _generate(owner, draft).json()["pdf_status"] == "failed"
    response = _download(owner, draft, 1, "pdf")
    assert response.status_code == 409
    assert response.json()["detail"] == "Version 1 has no PDF"
    storage_service.delete_file(_latest(draft).docx_path)
    missing = _download(owner, draft, 1, "docx")
    assert missing.status_code == 410
    row = _versions(owner, draft).json()[0]
    assert (row["has_docx"], row["docx_url"], row["files_pruned"]) == (False, None, False)


def test_retention_prunes_old_version_files(owner, draft, monkeypatch):
    monkeypatch.setattr(settings, "generated_versions_keep_files", 1)
    _generate(owner, draft)
    v1 = _latest(draft)
    _put(owner, draft, [{"field_name": "title", "value": "Later"}])
    _generate(owner, draft)
    assert not storage_service.exists(v1.docx_path)
    assert not storage_service.exists(v1.pdf_path)
    gone = _download(owner, draft, 1, "docx")
    assert gone.status_code == 410
    assert "removed to save space" in gone.json()["detail"]
    assert _download(owner, draft, 1, "pdf").status_code == 410
    rows = {row["version"]: row for row in _versions(owner, draft).json()}
    assert (rows[1]["files_pruned"], rows[1]["has_docx"], rows[1]["has_pdf"], rows[1]["docx_url"]) == (True, False, False, None)
    assert rows[2]["files_pruned"] is False and rows[2]["has_docx"] is True
    detail = client.get(f"{DOCS}/{draft}/versions/1", headers=_auth(owner))
    assert detail.status_code == 200 and detail.json()["files_pruned"] is True


def test_restore_then_generate_creates_v3(owner, draft):
    _generate(owner, draft)
    _put(owner, draft, [{"field_name": "title", "value": "Changed"}, {"field_name": "notes", "value": "added later"},
                        {"field_name": "discussion", "value": "plain now"}])
    _generate(owner, draft)
    restored = client.post(f"{DOCS}/{draft}/versions/1/restore", headers=_auth(owner))
    assert restored.status_code == 200, restored.text
    values = {v["field_name"]: v for v in restored.json()["values"]}
    assert set(values) == {"title", "discussion"}  # "notes" was not in v1
    assert values["title"]["value"] == "Budget"
    assert values["discussion"]["value_format"] == "rich"
    assert values["discussion"]["rich_content"] == validate_rich_content(SHARED_DOC)
    assert (values["discussion"]["ai_modified"], values["discussion"]["ai_action"]) == (True, "formal")
    assert restored.json()["generated"]["is_stale"] is True
    third = _generate(owner, draft).json()
    assert third["version"] == 3
    rows = _versions(owner, draft).json()
    assert rows[0]["changed_fields"] == ["discussion", "notes", "title"]
    detail = client.get(f"{DOCS}/{draft}/versions/3", headers=_auth(owner)).json()
    detail_v1 = client.get(f"{DOCS}/{draft}/versions/1", headers=_auth(owner)).json()
    assert detail["sections"] == detail_v1["sections"]


def test_restore_rules(owner, draft):
    assert client.post(f"{DOCS}/{draft}/versions/1/restore", headers=_auth(owner)).status_code == 404
    _generate(owner, draft)
    with TestSession() as session:
        session.get(Document, draft).status = "submitted"
        session.commit()
    response = client.post(f"{DOCS}/{draft}/versions/1/restore", headers=_auth(owner))
    assert response.status_code == 409
    # Viewing and downloading a submitted document's versions still works.
    assert _versions(owner, draft).status_code == 200
    assert _download(owner, draft, 1, "docx").status_code == 200


def test_version_routes_are_owner_only(owner, stranger, draft):
    _generate(owner, draft)
    for method, url in (
        ("get", f"{DOCS}/{draft}/versions"),
        ("get", f"{DOCS}/{draft}/versions/1"),
        ("get", f"{DOCS}/{draft}/versions/1/download/docx"),
        ("get", f"{DOCS}/{draft}/versions/1/download/pdf"),
        ("post", f"{DOCS}/{draft}/versions/1/restore"),
    ):
        assert getattr(client, method)(url, headers=_auth(stranger)).status_code == 403, url
        assert getattr(client, method)(url).status_code == 401, url
    assert client.get(f"{DOCS}/999999/versions", headers=_auth(owner)).status_code == 404
    assert client.get(f"{DOCS}/{draft}/versions/7", headers=_auth(owner)).status_code == 404
    assert _versions(owner, _create(owner, _make_template())).json() == []


def test_version_detail_keeps_fields_removed_from_template(owner, draft):
    _generate(owner, draft)
    with TestSession() as session:
        template_id = session.get(Document, draft).template_id
        from app.models.template_field import TemplateField

        session.execute(
            TemplateField.__table__.delete().where(
                TemplateField.template_id == template_id, TemplateField.field_name == "discussion"
            )
        )
        session.commit()
    detail = client.get(f"{DOCS}/{draft}/versions/1", headers=_auth(owner)).json()
    fields = {f["field_name"]: f for s in detail["sections"] for f in s["fields"]}
    assert fields["discussion"]["display_value"] == SHARED_PLAIN
    assert fields["discussion"]["field_type"] == "textarea"
    assert fields["discussion"]["label"] == "discussion"


# ── Preview ───────────────────────────────────────────────────────


def test_preview_carries_rich_ai_and_version_fields(owner, draft):
    preview = client.get(f"{DOCS}/{draft}/preview", headers=_auth(owner)).json()
    fields = {f["field_name"]: f for s in preview["sections"] for f in s["fields"]}
    assert fields["discussion"]["value_format"] == "rich"
    assert fields["discussion"]["rich_content"] == validate_rich_content(SHARED_DOC)
    assert fields["discussion"]["display_value"] == SHARED_PLAIN
    assert (fields["discussion"]["ai_modified"], fields["discussion"]["ai_action"]) == (True, "formal")
    assert (fields["title"]["value_format"], fields["title"]["ai_modified"]) == ("plain", False)
    assert fields["notes"]["rich_content"] is None
    assert preview["generation"] is None
    _generate(owner, draft)
    _generate(owner, draft)
    generation = client.get(f"{DOCS}/{draft}/preview", headers=_auth(owner)).json()["generation"]
    assert (generation["version"], generation["versions_count"]) == (2, 2)


def test_plain_projection_matches_docx_text(owner):
    """The preview / saved plain value is exactly the text in the DOCX."""
    document_id = _create(owner, _make_template(["{{ discussion }}"], [("discussion", "textarea", False, None)]))
    _put(owner, document_id, [_rich("discussion", SHARED_DOC)])
    _generate(owner, document_id)
    texts = _paragraph_texts(storage_service.read_bytes(_latest(document_id).docx_path))
    assert texts == [to_plain_text(validate_rich_content(SHARED_DOC))]


# ── Cleanup ───────────────────────────────────────────────────────


def test_document_delete_removes_every_version_file(owner, draft):
    _generate(owner, draft)
    _generate(owner, draft)
    with TestSession() as session:
        paths = [
            path
            for row in session.execute(select(GeneratedFile).where(GeneratedFile.document_id == draft)).scalars()
            for path in (row.docx_path, row.pdf_path)
        ]
    assert len(paths) == 4 and all(storage_service.exists(p) for p in paths)
    with TestSession() as session:
        assert delete_document(session, draft) is True
    assert not any(storage_service.exists(p) for p in paths)
