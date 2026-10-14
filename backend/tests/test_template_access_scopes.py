"""
Organization / department scoped templates with role-restricted audiences,
matched by organization / department ids.

- View rule vs library SQL filter parity across a users x templates matrix.
- can_manage.
- Publish rules (resolve_publish_access) and their HTTP mapping.
- POST /templates/upload with scope ids (and the name fallback),
  PATCH /templates/{id}/access, GET /templates/access-options, and
  can_manage-gated modifying endpoints.
"""

import io
import itertools
import json
import os

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests")

from app.core.security import create_access_token
from app.crud.template_crud import get_library_templates
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.organization import Department, Organization
from app.models.template import Template
from app.models.user import User
from app.services.template_access import (
    PublishRuleError,
    can_manage,
    resolve_publish_access,
    user_can_view_template,
)

test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestSession = sessionmaker(bind=test_engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)

UNI = "State University"
OTHER = "Other College"

# key: organization name
ORGANIZATIONS = {"uni": UNI, "other": OTHER}
# key: (organization key, department name). Both organizations have a
# "Computer Science" department: same name, different rows.
DEPARTMENTS = {
    "uni_cs": ("uni", "Computer Science"),
    "uni_me": ("uni", "Mechanical"),
    "uni_phys": ("uni", "Physics"),
    "other_cs": ("other", "Computer Science"),
    "other_law": ("other", "Law"),
}

# key: (role, organization key, department key)
USERS = {
    "super": ("super_admin", None, None),
    "uni_org_admin": ("org_admin", "uni", None),
    "other_org_admin": ("org_admin", "other", None),
    "cs_dept_admin": ("department_admin", "uni", "uni_cs"),
    "me_dept_admin": ("department_admin", "uni", "uni_me"),
    "cs_faculty": ("faculty", "uni", "uni_cs"),
    "cs_student": ("student", "uni", "uni_cs"),
    "cs_staff": ("normal_user", "uni", "uni_cs"),
    "me_faculty": ("faculty", "uni", "uni_me"),
    "other_cs_faculty": ("faculty", "other", "other_cs"),
    "no_org_student": ("student", None, None),
    "approver": ("approver", "uni", None),
}


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    old = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


@pytest.fixture(scope="module")
def directory(setup_db) -> dict[str, int]:
    """Organization and department ids by key."""
    ids: dict[str, int] = {}
    with TestSession() as session:
        for key, name in ORGANIZATIONS.items():
            organization = Organization(name=name)
            session.add(organization)
            session.flush()
            ids[key] = organization.id
        for key, (org_key, name) in DEPARTMENTS.items():
            department = Department(organization_id=ids[org_key], name=name)
            session.add(department)
            session.flush()
            ids[key] = department.id
        session.commit()
    return ids


@pytest.fixture(scope="module")
def users(directory) -> dict[str, int]:
    ids = {}
    with TestSession() as session:
        for key, (role, org, dept) in USERS.items():
            user = User(
                email=f"{key}@access.test",
                full_name=key,
                role=role,
                organization_id=directory[org] if org else None,
                department_id=directory[dept] if dept else None,
            )
            session.add(user)
            session.flush()
            ids[key] = user.id
        session.commit()
    return ids


def _token(users: dict[str, int], key: str) -> dict:
    return {"Authorization": f"Bearer {create_access_token(str(users[key]))}"}


# (name, owner key, visibility, scope org key, scope dept key, allowed_roles list)
TEMPLATE_SPECS = [
    ("private", "cs_faculty", "private", None, None, None),
    ("public", "cs_faculty", "public", None, None, None),
    ("public_faculty_only", "uni_org_admin", "public", None, None, ["faculty"]),
    ("uni_letterhead", "uni_org_admin", "organization", "uni", None, None),
    ("uni_letterhead_2", "super", "organization", "uni", None, None),
    ("uni_staff_only", "uni_org_admin", "organization", "uni", None, ["faculty", "normal_user"]),
    ("other_letterhead", "other_org_admin", "organization", "other", None, None),
    ("cs_dept", "cs_dept_admin", "department", "uni", "uni_cs", None),
    (
        "cs_dept_no_students",
        "cs_dept_admin",
        "department",
        "uni",
        "uni_cs",
        ["faculty", "normal_user", "department_admin"],
    ),
    # Inconsistent legacy row: department set, organization not.
    ("cs_dept_org_unset", "super", "department", None, "uni_cs", None),
    # Same department NAME as uni_cs but another organization's row.
    ("other_cs_dept", "super", "department", "other", "other_cs", None),
    ("me_dept", "me_dept_admin", "department", "uni", "uni_me", None),
    ("group_faculty", "cs_faculty", "group", None, None, None),
    ("legacy_org_unscoped", "cs_faculty", "organization", None, None, None),
    # Backfill leftover: department string without organization -> no ids.
    ("dept_unlinked", "super", "department", None, None, None),
    ("malformed_roles", "uni_org_admin", "public", None, None, "not-json"),
    ("empty_roles", "uni_org_admin", "organization", "uni", None, "[]"),
]


@pytest.fixture(scope="module")
def templates(users, directory) -> dict[str, int]:
    ids = {}
    with TestSession() as session:
        for name, owner, visibility, org, dept, roles in TEMPLATE_SPECS:
            if isinstance(roles, list):
                roles = json.dumps(roles)
            template = Template(
                name=name,
                category="notice",
                visibility=visibility,
                uploaded_by=users[owner],
                scope_organization_id=directory[org] if org else None,
                scope_department_id=directory[dept] if dept else None,
                allowed_roles=roles,
            )
            session.add(template)
            session.flush()
            ids[name] = template.id
        session.commit()
    return ids


# ── parity: Python view rule == library SQL filter ───────────────────


def test_view_rule_matches_library_filter_for_every_user(users, templates):
    with TestSession() as session:
        all_templates = session.query(Template).all()
        assert len(all_templates) == len(TEMPLATE_SPECS)
        for key in USERS:
            user = session.get(User, users[key])
            expected = {t.name for t in all_templates if user_can_view_template(user, t)}
            listed, total = get_library_templates(session, user.id, limit=100)
            assert {t.name for t in listed} == expected, key
            assert total == len(expected), key


EXPECTED_VISIBLE = {
    "cs_student": {
        "public",
        "uni_letterhead",
        "uni_letterhead_2",
        "cs_dept",
        "cs_dept_org_unset",
        "empty_roles",
    },
    "cs_faculty": {
        "private",
        "public",
        "public_faculty_only",
        "uni_letterhead",
        "uni_letterhead_2",
        "uni_staff_only",
        "cs_dept",
        "cs_dept_no_students",
        "cs_dept_org_unset",
        "group_faculty",
        "legacy_org_unscoped",
        "empty_roles",
    },
    # Another organization's "Computer Science" is a different department.
    "other_cs_faculty": {
        "public",
        "public_faculty_only",
        "other_letterhead",
        "other_cs_dept",
        "group_faculty",
    },
    # Scope manager: sees restricted org + all dept templates of the org.
    "uni_org_admin": {
        "public",
        "public_faculty_only",
        "uni_letterhead",
        "uni_letterhead_2",
        "uni_staff_only",
        "cs_dept",
        "cs_dept_no_students",
        "me_dept",
        "malformed_roles",
        "empty_roles",
    },
    "other_org_admin": {"public", "other_letterhead", "other_cs_dept"},
    # Department manager sees its department even when restricted; not other depts.
    "me_dept_admin": {
        "public",
        "uni_letterhead",
        "uni_letterhead_2",
        "me_dept",
        "empty_roles",
    },
    "no_org_student": {"public"},
}


@pytest.mark.parametrize("key", sorted(EXPECTED_VISIBLE))
def test_view_rule_expected_sets(users, templates, key):
    with TestSession() as session:
        user = session.get(User, users[key])
        visible = {
            t.name for t in session.query(Template).all() if user_can_view_template(user, t)
        }
    assert visible == EXPECTED_VISIBLE[key]


def test_unlinked_department_template_is_owner_only(users, templates):
    with TestSession() as session:
        template = session.get(Template, templates["dept_unlinked"])
        viewers = {
            key
            for key in USERS
            if user_can_view_template(session.get(User, users[key]), template)
        }
    assert viewers == {"super"}


def test_super_admin_sees_everything(users, templates):
    with TestSession() as session:
        listed, total = get_library_templates(session, users["super"], limit=100)
    assert total == len(TEMPLATE_SPECS)


# ── can_manage ───────────────────────────────────────────────────────

EXPECTED_MANAGE = {
    "super": {name for name, *_ in TEMPLATE_SPECS},
    "uni_org_admin": {
        "public_faculty_only",
        "uni_letterhead",
        "uni_letterhead_2",
        "uni_staff_only",
        "cs_dept",
        "cs_dept_no_students",
        "me_dept",
        "malformed_roles",
        "empty_roles",
    },
    "other_org_admin": {"other_letterhead", "other_cs_dept"},
    "cs_dept_admin": {"cs_dept", "cs_dept_no_students", "cs_dept_org_unset"},
    "me_dept_admin": {"me_dept"},
    "cs_faculty": {"private", "public", "group_faculty", "legacy_org_unscoped"},
    "cs_student": set(),
}


@pytest.mark.parametrize("key", sorted(EXPECTED_MANAGE))
def test_can_manage(users, templates, key):
    with TestSession() as session:
        user = session.get(User, users[key])
        managed = {t.name for t in session.query(Template).all() if can_manage(user, t)}
    assert managed == EXPECTED_MANAGE[key]


def test_can_manage_implies_can_view(users, templates):
    with TestSession() as session:
        all_templates = session.query(Template).all()
        for key in USERS:
            user = session.get(User, users[key])
            for template in all_templates:
                if can_manage(user, template):
                    assert user_can_view_template(user, template), (key, template.name)


# ── publish rules ────────────────────────────────────────────────────


@pytest.fixture()
def db(directory):
    with TestSession() as session:
        yield session


def _user(directory, role, org="uni", dept=None):
    return User(
        email="x@x",
        full_name="x",
        role=role,
        organization_id=directory[org] if org else None,
        department_id=directory[dept] if dept else None,
    )


def _ids(resolved):
    return resolved.scope_organization_id, resolved.scope_department_id


@pytest.mark.parametrize("role", ["faculty", "student", "approver", "normal_user"])
@pytest.mark.parametrize("visibility", ["organization", "department"])
def test_non_admins_cannot_publish_scoped(db, directory, role, visibility):
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(
            db,
            _user(directory, role, dept="uni_cs"),
            visibility,
            directory["uni"],
            directory["uni_cs"],
            None,
        )
    assert exc.value.status_code == 403


def test_org_admin_scope_forced_to_own_org(db, directory):
    admin = _user(directory, "org_admin")
    resolved = resolve_publish_access(db, admin, "organization")
    assert _ids(resolved) == (directory["uni"], None)
    assert resolved.organization.name == UNI
    # Same org by id, or by name (fallback, case-insensitive) is accepted.
    assert _ids(
        resolve_publish_access(db, admin, "organization", directory["uni"])
    ) == (directory["uni"], None)
    assert _ids(
        resolve_publish_access(
            db, admin, "organization", scope_organization_name="  state UNIVERSITY "
        )
    ) == (directory["uni"], None)
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "organization", directory["other"])
    assert exc.value.status_code == 403
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "organization", scope_organization_name="Nope U")
    assert exc.value.status_code == 422
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, _user(directory, "org_admin", org=None), "organization")
    assert exc.value.status_code == 422


def test_org_admin_department_rules(db, directory):
    admin = _user(directory, "org_admin")
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "department")
    assert exc.value.status_code == 422
    resolved = resolve_publish_access(
        db, admin, "department", scope_department_id=directory["uni_phys"]
    )
    assert _ids(resolved) == (directory["uni"], directory["uni_phys"])
    by_name = resolve_publish_access(
        db, admin, "department", scope_department_name=" physics "
    )
    assert _ids(by_name) == (directory["uni"], directory["uni_phys"])
    # Another organization's department.
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(
            db, admin, "department", scope_department_id=directory["other_law"]
        )
    assert exc.value.status_code == 403
    # Unknown department.
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "department", scope_department_id=999_999)
    assert exc.value.status_code == 422
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "department", scope_department_name="Law")
    assert exc.value.status_code == 422


def test_department_admin_rules(db, directory):
    admin = _user(directory, "department_admin", dept="uni_cs")
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "organization")
    assert exc.value.status_code == 403
    resolved = resolve_publish_access(db, admin, "department")
    assert _ids(resolved) == (directory["uni"], directory["uni_cs"])
    # Explicitly naming their own department / organization is fine.
    assert _ids(
        resolve_publish_access(
            db, admin, "department", directory["uni"], directory["uni_cs"]
        )
    ) == (directory["uni"], directory["uni_cs"])
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(
            db, admin, "department", scope_department_id=directory["uni_me"]
        )
    assert exc.value.status_code == 403
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "department", directory["other"])
    assert exc.value.status_code == 403
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(
            db, _user(directory, "department_admin", dept=None), "department"
        )
    assert exc.value.status_code == 422


def test_super_admin_must_name_scope(db, directory):
    admin = _user(directory, "super_admin", org=None)
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "organization")
    assert exc.value.status_code == 422
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "department", directory["other"])
    assert exc.value.status_code == 422
    resolved = resolve_publish_access(
        db, admin, "department", directory["other"], directory["other_law"]
    )
    assert _ids(resolved) == (directory["other"], directory["other_law"])
    assert (resolved.organization.name, resolved.department.name) == (OTHER, "Law")
    # The department alone implies its organization.
    inferred = resolve_publish_access(
        db, admin, "department", scope_department_id=directory["other_law"]
    )
    assert _ids(inferred) == (directory["other"], directory["other_law"])
    # Department of another organization than the one named.
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(
            db, admin, "department", directory["uni"], directory["other_law"]
        )
    assert exc.value.status_code == 422
    with pytest.raises(PublishRuleError) as exc:
        resolve_publish_access(db, admin, "organization", 999_999)
    assert exc.value.status_code == 422
    # Organization visibility drops a department.
    org_only = resolve_publish_access(
        db, admin, "organization", directory["uni"], directory["uni_cs"]
    )
    assert _ids(org_only) == (directory["uni"], None)


def test_allowed_roles_validation(db, directory):
    admin = _user(directory, "org_admin")
    resolved = resolve_publish_access(
        db, admin, "organization", None, None, ["student", "faculty", "faculty"]
    )
    assert json.loads(resolved.allowed_roles) == ["faculty", "student"]
    assert (
        resolve_publish_access(db, admin, "organization", None, None, []).allowed_roles
        is None
    )
    for bad in (["super_admin"], ["wizard"]):
        with pytest.raises(PublishRuleError) as exc:
            resolve_publish_access(db, admin, "organization", None, None, bad)
        assert exc.value.status_code == 422


def test_unscoped_visibilities_clear_scope(db, directory):
    faculty = _user(directory, "faculty", dept="uni_cs")
    resolved = resolve_publish_access(
        db, faculty, "public", directory["uni"], directory["uni_cs"], ["faculty"]
    )
    assert _ids(resolved) == (None, None)
    assert json.loads(resolved.allowed_roles) == ["faculty"]
    private = resolve_publish_access(db, faculty, "private", None, None, ["faculty"])
    assert private.allowed_roles is None


# ── HTTP: upload / access / access-options ───────────────────────────


def _docx() -> bytes:
    document = Document()
    document.add_paragraph("Dear {{recipient_name}},")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(headers, **data):
    payload = {"name": "Scoped", "category": "notice", "visibility": "private", **data}
    return client.post(
        "/api/v1/templates/upload",
        files={
            "file": (
                "scoped.docx",
                io.BytesIO(_docx()),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data=payload,
        headers=headers,
    )


def test_upload_department_template_by_department_admin(users, directory):
    response = _upload(
        _token(users, "cs_dept_admin"),
        visibility="department",
        allowed_roles="faculty, normal_user",
    )
    assert response.status_code == 201, response.json()
    body = response.json()
    assert body["scope_organization"] == UNI
    assert body["scope_department"] == "Computer Science"
    assert body["scope_organization_id"] == directory["uni"]
    assert body["scope_department_id"] == directory["uni_cs"]
    assert body["allowed_roles"] == ["faculty", "normal_user"]
    assert body["can_manage"] is True

    # The legacy name columns are written as copies of the linked names.
    with TestSession() as session:
        stored = session.get(Template, body["id"])
        assert (stored.scope_organization, stored.scope_department) == (
            UNI,
            "Computer Science",
        )

    # Faculty in the department can open it; the student cannot.
    template_id = body["id"]
    faculty = client.get(f"/api/v1/templates/{template_id}", headers=_token(users, "cs_faculty"))
    assert faculty.status_code == 200
    assert faculty.json()["can_manage"] is False
    student = client.get(f"/api/v1/templates/{template_id}", headers=_token(users, "cs_student"))
    assert student.status_code == 403
    # ...and it is not in the student's library.
    library = client.get("/api/v1/templates/library", headers=_token(users, "cs_student"))
    assert template_id not in {t["id"] for t in library.json()["templates"]}
    # Same department name in another organization does not match.
    other = client.get(
        f"/api/v1/templates/{template_id}", headers=_token(users, "other_cs_faculty")
    )
    assert other.status_code == 403


def test_upload_rejections(users, directory):
    assert _upload(_token(users, "cs_faculty"), visibility="organization").status_code == 403
    response = _upload(_token(users, "cs_dept_admin"), visibility="organization")
    assert response.status_code == 403
    assert "department" in response.json()["detail"].lower()
    assert _upload(_token(users, "super"), visibility="organization").status_code == 422
    assert (
        _upload(_token(users, "uni_org_admin"), visibility="organization", allowed_roles="god")
        .status_code
        == 422
    )
    assert (
        _upload(
            _token(users, "uni_org_admin"),
            visibility="organization",
            scope_organization_id=str(directory["other"]),
        ).status_code
        == 403
    )
    assert (
        _upload(
            _token(users, "super"),
            visibility="department",
            scope_organization_id=str(directory["uni"]),
            scope_department_id=str(directory["other_law"]),
        ).status_code
        == 422
    )


def test_upload_super_admin_scoped(users, directory):
    response = _upload(
        _token(users, "super"),
        visibility="organization",
        scope_organization_id=str(directory["other"]),
        allowed_roles="",
    )
    assert response.status_code == 201
    assert response.json()["scope_organization"] == OTHER
    assert response.json()["scope_organization_id"] == directory["other"]
    assert response.json()["allowed_roles"] is None

    # Name fallback for older clients.
    by_name = _upload(
        _token(users, "super"),
        visibility="department",
        scope_organization="other college",
        scope_department="LAW",
    )
    assert by_name.status_code == 201, by_name.json()
    assert by_name.json()["scope_department_id"] == directory["other_law"]
    assert by_name.json()["scope_department"] == "Law"
    unknown = _upload(
        _token(users, "super"), visibility="organization", scope_organization="Nowhere"
    )
    assert unknown.status_code == 422


def test_library_items_carry_audience(users, directory):
    response = client.get("/api/v1/templates/library?limit=100", headers=_token(users, "uni_org_admin"))
    assert response.status_code == 200
    items = {t["name"]: t for t in response.json()["templates"]}
    staff_only = items["uni_staff_only"]
    assert staff_only["scope_organization"] == UNI
    assert staff_only["scope_organization_id"] == directory["uni"]
    assert staff_only["allowed_roles"] == ["faculty", "normal_user"]
    assert staff_only["can_manage"] is True
    assert items["public"]["can_manage"] is False
    assert items["me_dept"]["scope_department"] == "Mechanical"
    assert items["me_dept"]["scope_department_id"] == directory["uni_me"]


def test_patch_access(users, templates, directory):
    template_id = templates["cs_dept"]
    url = f"/api/v1/templates/{template_id}/access"

    # Viewers who cannot manage get 403.
    denied = client.patch(url, json={"visibility": "private"}, headers=_token(users, "cs_faculty"))
    assert denied.status_code == 403

    # Department admin restricts it to faculty.
    response = client.patch(
        url,
        json={"visibility": "department", "allowed_roles": ["faculty"]},
        headers=_token(users, "cs_dept_admin"),
    )
    assert response.status_code == 200, response.json()
    assert response.json()["allowed_roles"] == ["faculty"]
    assert response.json()["scope_department"] == "Computer Science"
    assert response.json()["scope_department_id"] == directory["uni_cs"]
    assert client.get(f"/api/v1/templates/{template_id}", headers=_token(users, "cs_staff")).status_code == 403

    # Department admin cannot widen to the organization or move it.
    widen = client.patch(url, json={"visibility": "organization"}, headers=_token(users, "cs_dept_admin"))
    assert widen.status_code == 403
    move = client.patch(
        url,
        json={"visibility": "department", "scope_department_id": directory["uni_me"]},
        headers=_token(users, "cs_dept_admin"),
    )
    assert move.status_code == 403

    # The org admin (scope manager) can, and clearing roles reopens it.
    response = client.patch(
        url,
        json={"visibility": "organization", "allowed_roles": []},
        headers=_token(users, "uni_org_admin"),
    )
    assert response.status_code == 200
    body = response.json()
    assert body["visibility"] == "organization"
    assert body["scope_organization"] == UNI
    assert body["scope_organization_id"] == directory["uni"]
    assert body["scope_department"] is None
    assert body["scope_department_id"] is None
    assert body["allowed_roles"] is None
    assert client.get(f"/api/v1/templates/{template_id}", headers=_token(users, "cs_staff")).status_code == 200

    # ...and back to a department of their organization, by id.
    response = client.patch(
        url,
        json={"visibility": "department", "scope_department_id": directory["uni_cs"]},
        headers=_token(users, "uni_org_admin"),
    )
    assert response.status_code == 200
    assert response.json()["scope_department_id"] == directory["uni_cs"]

    invalid = client.patch(url, json={"visibility": "everyone"}, headers=_token(users, "super"))
    assert invalid.status_code == 422
    mismatch = client.patch(
        url,
        json={
            "visibility": "department",
            "scope_organization_id": directory["uni"],
            "scope_department_id": directory["other_cs"],
        },
        headers=_token(users, "super"),
    )
    assert mismatch.status_code == 422
    missing = client.patch("/api/v1/templates/999999/access", json={"visibility": "private"}, headers=_token(users, "super"))
    assert missing.status_code == 404


def test_access_options(users, directory):
    response = client.get("/api/v1/templates/access-options", headers=_token(users, "super"))
    assert response.status_code == 200
    body = response.json()
    role_values = [role["value"] for role in body["roles"]]
    assert "super_admin" not in role_values
    assert {"faculty", "student", "normal_user"} <= set(role_values)
    assert all(role["label"] for role in body["roles"])
    assert body["organizations"] == [
        {"id": directory["other"], "name": OTHER},
        {"id": directory["uni"], "name": UNI},
    ]
    departments = {
        (d["organization_id"], d["name"]): d["id"] for d in body["departments"]
    }
    assert departments[(directory["uni"], "Computer Science")] == directory["uni_cs"]
    assert departments[(directory["other"], "Computer Science")] == directory["other_cs"]
    assert len(departments) == len(DEPARTMENTS)
    assert body["can_publish"] == {"organization": True, "department": True}
    assert body["defaults"] == {"organization_id": None, "department_id": None}

    dept_admin = client.get(
        "/api/v1/templates/access-options", headers=_token(users, "cs_dept_admin")
    ).json()
    assert dept_admin["can_publish"] == {"organization": False, "department": True}
    assert dept_admin["defaults"] == {
        "organization_id": directory["uni"],
        "department_id": directory["uni_cs"],
    }
    assert dept_admin["organizations"] == [{"id": directory["uni"], "name": UNI}]
    assert {d["organization_id"] for d in dept_admin["departments"]} == {directory["uni"]}

    faculty = client.get(
        "/api/v1/templates/access-options", headers=_token(users, "cs_faculty")
    ).json()
    assert faculty["can_publish"] == {"organization": False, "department": False}

    no_org = client.get(
        "/api/v1/templates/access-options", headers=_token(users, "no_org_student")
    ).json()
    assert no_org["organizations"] == [] and no_org["departments"] == []


def test_access_options_route_not_shadowed_by_template_id(users):
    response = client.get("/api/v1/templates/access-options", headers=_token(users, "cs_student"))
    assert response.status_code == 200


def test_scope_manager_can_use_modifying_endpoints(users, templates):
    template_id = templates["me_dept"]
    field = {"field_name": "dean_name", "field_label": "Dean", "field_type": "text"}
    # The org admin manages the Mechanical department template.
    created = client.post(
        f"/api/v1/templates/{template_id}/fields",
        json=field,
        headers=_token(users, "uni_org_admin"),
    )
    assert created.status_code == 201, created.json()
    # A faculty member of that department can view but not modify.
    assert (
        client.get(f"/api/v1/templates/{template_id}/fields", headers=_token(users, "me_faculty")).status_code
        == 200
    )
    denied = client.post(
        f"/api/v1/templates/{template_id}/fields",
        json={**field, "field_name": "other"},
        headers=_token(users, "me_faculty"),
    )
    assert denied.status_code == 403
    # Another department's admin cannot either.
    assert (
        client.post(
            f"/api/v1/templates/{template_id}/detect-placeholders",
            headers=_token(users, "cs_dept_admin"),
        ).status_code
        == 403
    )


def test_delete_stays_super_admin_only(users, templates):
    response = client.delete(
        f"/api/v1/templates/{templates['other_letterhead']}",
        headers=_token(users, "other_org_admin"),
    )
    assert response.status_code == 403


def test_matrix_is_meaningful():
    # Guard: the parity matrix exercises every visibility and role restriction.
    visibilities = {spec[2] for spec in TEMPLATE_SPECS}
    assert visibilities == {"private", "public", "organization", "department", "group"}
    assert any(isinstance(spec[5], list) for spec in TEMPLATE_SPECS)
    assert len(list(itertools.product(USERS, TEMPLATE_SPECS))) >= 150
