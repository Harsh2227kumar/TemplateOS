"""
V1.7 Phase 2.1 data layer (Member 3): template revisions, personal copies,
field changes, editor sessions, manual-edit versions.

Coverage:
- add_revision: placeholders added / removed vs the fields, blank suggestions.
- field_changes: new placeholder (type inferred), removed field, blank
  "Date of meeting: ______", unparseable template.
- copy_template_for_user: new stored files with equal bytes, fields copied,
  private + parent, original untouched, name numbering, customize revision.
- editor sessions: same user reuses, other user locked, expiry frees the
  lock, unique keys, status helpers, error truncation.
- create_version(source="manual_edit"): snapshot from the base version.
- Cascades: template delete removes revisions + sessions and queues their
  files; deleting a base version keeps the manual edit.
- save_processed keeps the superseded file (revisions reference it).

Database: in-memory SQLite with PRAGMA foreign_keys=ON; storage in tmp.
"""

import io
import json
import os
from datetime import timedelta

import pytest
from docx import Document as DocxDocument
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault("JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests")

import app.models  # noqa: F401,E402
from app.api.v1.endpoints._template_common import save_processed  # noqa: E402
from app.core.clock import utc_now_naive  # noqa: E402
from app.core.config import Settings  # noqa: E402
from app.crud import editor_session_crud as sessions  # noqa: E402
from app.crud.document_version_crud import create_version, parse_snapshot  # noqa: E402
from app.crud.template_copy_crud import copy_template_for_user  # noqa: E402
from app.crud.template_revision_crud import (  # noqa: E402
    add_revision,
    field_changes,
    get_revision,
    list_revisions,
)
from app.db.base import Base  # noqa: E402
from app.models.document import Document  # noqa: E402
from app.models.document_version import DocumentVersion  # noqa: E402
from app.models.editor_session import EditorSession  # noqa: E402
from app.models.template import Template  # noqa: E402
from app.models.template_field import TemplateField  # noqa: E402
from app.models.template_revision import TemplateRevision  # noqa: E402
from app.models.user import User  # noqa: E402
from app.schemas.document_version import DocumentVersionRead  # noqa: E402
from app.schemas.template import TemplateResponse  # noqa: E402
from app.services.storage_service import storage_service  # noqa: E402

engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)


@event.listens_for(engine, "connect")
def _enable_foreign_keys(dbapi_connection, _record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


TestSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)


@pytest.fixture(scope="module", autouse=True)
def setup_db(tmp_path_factory):
    old_root = storage_service.root_path
    storage_service.root_path = tmp_path_factory.mktemp("editor-storage").resolve()
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    yield
    Base.metadata.drop_all(bind=engine)
    storage_service.root_path = old_root


def docx(*paragraphs: str) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


_n = {"value": 0}


@pytest.fixture
def db():
    with TestSession() as session:
        yield session


def _user(db, role="faculty") -> User:
    _n["value"] += 1
    user = User(email=f"editor{_n['value']}@example.com", full_name=f"User {_n['value']}", role=role)
    db.add(user)
    db.flush()
    return user


@pytest.fixture
def admin_template(db) -> Template:
    admin = _user(db, "org_admin")
    original = storage_service.save_bytes("templates_original", "mom.docx", docx("Title: ____")).path
    processed = storage_service.save_bytes("templates_processed", "mom.docx", docx("Title: {{ title }}")).path
    template = Template(
        name="Minutes of Meeting",
        category="mom",
        visibility="public",
        status="active",
        uploaded_by=admin.id,
        original_file_path=original,
        processed_file_path=processed,
        original_filename="mom.docx",
        version=2,
    )
    db.add(template)
    db.flush()
    db.add(TemplateField(template_id=template.id, field_name="title", field_label="Title",
                         field_type="text", is_required=True, display_order=0, ai_enabled=True))
    db.commit()
    return template


# ── Revisions and field changes ──────────────────────────────────


def test_add_revision_records_placeholder_changes(db, admin_template):
    path = storage_service.save_bytes(
        "templates_revisions", "mom.docx",
        docx("Date of meeting: {{ meeting_date }}", "Venue: ________________", "No title any more"),
    ).path
    admin_template.version = 3
    revision = add_revision(db, admin_template, source="manual_edit", file_path=path,
                            created_by=admin_template.uploaded_by, note="x" * 300)
    db.commit()
    assert (revision.revision, revision.source) == (3, "manual_edit")
    assert revision.added_names == ["meeting_date"]
    assert revision.removed_names == ["title"]
    assert revision.blank_suggestions >= 1
    assert len(revision.note) == 255
    assert [r.revision for r in list_revisions(db, admin_template.id)] == [3]
    assert get_revision(db, admin_template.id, 3).id == revision.id
    assert get_revision(db, admin_template.id, 9) is None
    with pytest.raises(ValueError):
        TemplateRevision(template_id=1, revision=1, source="magic", file_path="x")


def test_field_changes(db, admin_template):
    unchanged = field_changes(db, admin_template)
    assert (unchanged.added, unchanged.removed, unchanged.renderable) == ([], [], True)
    assert unchanged.has_changes is False

    admin_template.processed_file_path = storage_service.save_bytes(
        "templates_revisions", "mom.docx",
        docx("Date of meeting: {{ meeting_date }}", "Attendees: ____________"),
    ).path
    changes = field_changes(db, admin_template)
    assert [(a.key, a.suggested_type) for a in changes.added] == [("meeting_date", "date")]
    assert changes.added[0].suggested_label == "Meeting Date"
    assert changes.removed == ["title"]
    assert changes.blank_suggestions >= 1 and changes.has_changes

    admin_template.processed_file_path = storage_service.save_bytes(
        "templates_revisions", "mom.docx", docx("Title: {{ title", "{% if %}")
    ).path
    broken = field_changes(db, admin_template)
    assert broken.renderable is False and broken.render_error
    assert broken.removed == []  # unknown when the file cannot be parsed

    admin_template.processed_file_path = "templates/revisions/missing.docx"
    assert field_changes(db, admin_template).render_error == "The template file is missing on the server"
    db.rollback()


# ── Personal copies ──────────────────────────────────────────────


def test_copy_template_for_user(db, admin_template):
    faculty = _user(db)
    copy = copy_template_for_user(db, admin_template, faculty)
    db.commit()
    assert copy.id != admin_template.id
    assert (copy.uploaded_by, copy.visibility, copy.parent_template_id) == (faculty.id, "private", admin_template.id)
    assert copy.name == "Minutes of Meeting (my copy)"
    assert copy.is_personal_copy and copy.parent_template_name == "Minutes of Meeting"
    assert (copy.status, copy.version) == ("active", 1)
    assert copy.processed_file_path != admin_template.processed_file_path
    assert copy.original_file_path != admin_template.original_file_path
    assert storage_service.read_bytes(copy.processed_file_path) == storage_service.read_bytes(admin_template.processed_file_path)
    fields = db.execute(select(TemplateField).where(TemplateField.template_id == copy.id)).scalars().all()
    assert [(f.field_name, f.field_label, f.field_type, f.is_required, f.ai_enabled) for f in fields] == [
        ("title", "Title", "text", True, True)
    ]
    (revision,) = list_revisions(db, copy.id)
    assert (revision.revision, revision.source, revision.file_path) == (1, "customize", copy.processed_file_path)
    # The original is untouched.
    db.refresh(admin_template)
    assert (admin_template.version, admin_template.visibility, admin_template.parent_template_id) == (2, "public", None)
    assert len(db.execute(select(TemplateField).where(TemplateField.template_id == admin_template.id)).scalars().all()) == 1
    # A second copy is numbered; schemas expose the copy fields.
    second = copy_template_for_user(db, admin_template, faculty)
    db.commit()
    assert second.name == "Minutes of Meeting (my copy 2)"
    response = TemplateResponse.model_validate(second)
    assert (response.parent_template_id, response.parent_template_name, response.is_personal_copy) == (
        admin_template.id, "Minutes of Meeting", True,
    )


def test_copy_of_archived_template_is_usable(db, admin_template):
    admin_template.status = "archived"
    admin_template.processed_file_path = None
    db.commit()
    copy = copy_template_for_user(db, admin_template, _user(db))
    db.commit()
    assert copy.status == "active"
    assert copy.processed_file_path is None
    assert list_revisions(db, copy.id)[0].file_path == copy.original_file_path


# ── Editor sessions ──────────────────────────────────────────────


def _open(db, user, template, ttl=60):
    return sessions.open_or_reuse_session(
        db, user=user, target_type="template", template=template,
        source_path=template.processed_file_path, ttl_minutes=ttl,
    )


def test_session_reuse_lock_and_expiry(db, admin_template):
    owner = db.get(User, admin_template.uploaded_by)
    first, reused = _open(db, owner, admin_template)
    db.commit()
    assert (first.status, reused, len(first.session_key)) == ("open", False, 32)
    again, reused = _open(db, owner, admin_template)
    assert (again.id, reused) == (first.id, True)

    other = _user(db, "org_admin")
    with pytest.raises(sessions.EditorLocked) as locked:
        _open(db, other, admin_template)
    assert locked.value.user_name == owner.full_name
    assert "Being edited by" in str(locked.value)

    first.expires_at = utc_now_naive() - timedelta(minutes=1)
    db.commit()
    taken, reused = _open(db, other, admin_template)
    db.commit()
    db.refresh(first)
    assert (first.status, reused, taken.user_id) == ("expired", False, other.id)
    assert taken.session_key != first.session_key
    assert sessions.get_session_by_key(db, taken.session_key).id == taken.id
    assert sessions.get_session_by_key(db, "nope") is None
    assert sessions.get_session_by_key(db, "x" * 100) is None
    assert sessions.latest_session_for_target(db, "template", template_id=admin_template.id).id == taken.id


def test_session_status_helpers(db, admin_template):
    owner = db.get(User, admin_template.uploaded_by)
    session, _ = _open(db, owner, admin_template)
    sessions.mark_saving(session)
    assert session.is_active
    sessions.mark_saved(session, result_revision=3)
    assert (session.status, session.result_revision, session.is_active) == ("saved", 3, False)
    sessions.mark_failed(session, "x" * 800)
    assert session.status == "failed" and len(session.error) == 500 and session.error.endswith("...")
    sessions.mark_closed(session)
    db.commit()
    # A closed session frees the lock for everyone.
    fresh, reused = _open(db, _user(db, "org_admin"), admin_template)
    assert reused is False and fresh.id != session.id
    with pytest.raises(ValueError):
        EditorSession(session_key="k", target_type="spreadsheet", source_path="x", user_id=1,
                      expires_at=utc_now_naive())
    db.rollback()


def test_document_session_other_version_closes_previous(db, admin_template):
    owner = db.get(User, admin_template.uploaded_by)
    document = Document(template_id=admin_template.id, created_by=owner.id, name="Draft")
    db.add(document)
    db.flush()
    v1 = create_version(db, document, created_by=owner.id, template_version=2, snapshot=[])
    v2 = create_version(db, document, created_by=owner.id, template_version=2, snapshot=[])
    db.commit()
    s1, _ = sessions.open_or_reuse_session(db, user=owner, target_type="document", document=document,
                                           base_version=v1, source_path="a.docx", ttl_minutes=60)
    s2, reused = sessions.open_or_reuse_session(db, user=owner, target_type="document", document=document,
                                                base_version=v2, source_path="b.docx", ttl_minutes=60)
    db.commit()
    db.refresh(s1)
    assert (s1.status, reused, s2.base_version_id) == ("closed", False, v2.id)


# ── Manual-edit versions ─────────────────────────────────────────


def test_manual_edit_version_copies_base_snapshot(db, admin_template):
    owner = db.get(User, admin_template.uploaded_by)
    document = Document(template_id=admin_template.id, created_by=owner.id, name="Draft")
    db.add(document)
    db.flush()
    snapshot = [{"field_name": "title", "value": "Budget", "value_format": "plain",
                 "rich_content": None, "ai_modified": False, "ai_action": None}]
    base = create_version(db, document, created_by=owner.id, template_version=2, snapshot=snapshot)
    edited = create_version(db, document, created_by=owner.id, template_version=2,
                            source="manual_edit", base_version=base)
    db.commit()
    assert (edited.version, edited.source, edited.base_version_id) == (2, "manual_edit", base.id)
    assert parse_snapshot(edited) == parse_snapshot(base) == snapshot
    assert base.source == "generated" and edited.base_version.id == base.id
    with pytest.raises(ValueError):
        DocumentVersion(document_id=document.id, version=9, values_snapshot="[]", source="imported")
    read = DocumentVersionRead(
        version=2, created_at=edited.created_at, created_by_name="x", template_version=2, pdf_status=None,
        has_docx=False, has_pdf=False, files_pruned=False, changed_fields=[], is_latest=True,
        source="manual_edit", base_version=1,
    )
    assert read.source == "manual_edit"
    db.delete(base)
    db.commit()
    db.refresh(edited)
    assert edited.base_version_id is None


# ── Storage rules ────────────────────────────────────────────────


def test_save_processed_keeps_superseded_file(db, admin_template):
    previous = admin_template.processed_file_path
    new_path = save_processed(db, admin_template, docx("Title: {{ title }}", "Venue: {{ venue }}"))
    assert new_path != previous
    assert storage_service.exists(previous) and storage_service.exists(new_path)


def test_template_delete_removes_revisions_sessions_and_files(db, admin_template):
    owner = db.get(User, admin_template.uploaded_by)
    path = storage_service.save_bytes("templates_revisions", "mom.docx", docx("{{ title }}")).path
    admin_template.version = 3
    add_revision(db, admin_template, source="manual_edit", file_path=path, created_by=owner.id)
    _open(db, owner, admin_template)
    db.commit()
    template_id = admin_template.id
    db.delete(admin_template)
    db.commit()
    assert db.execute(select(TemplateRevision).where(TemplateRevision.template_id == template_id)).first() is None
    assert db.execute(select(EditorSession).where(EditorSession.template_id == template_id)).first() is None
    assert not storage_service.exists(path)


def test_editor_settings_validation():
    base = {"DATABASE_URL": "sqlite://", "JWT_SECRET_KEY": "x" * 40}
    off = Settings(**base)
    assert (off.onlyoffice_enabled, off.editor_session_ttl_minutes, off.editor_max_file_mb) == (False, 240, 20)
    on = Settings(**base, ONLYOFFICE_ENABLED=True, ONLYOFFICE_PUBLIC_URL="http://localhost:8082/",
                  ONLYOFFICE_JWT_SECRET="s" * 32)
    assert on.onlyoffice_public_base == "http://localhost:8082"
    assert on.onlyoffice_internal_base == "http://localhost:8082"
    with pytest.raises(ValueError):
        Settings(**base, ONLYOFFICE_ENABLED=True, ONLYOFFICE_JWT_SECRET="s" * 32)
    with pytest.raises(ValueError):
        Settings(**base, ONLYOFFICE_ENABLED=True, ONLYOFFICE_PUBLIC_URL="http://x", ONLYOFFICE_JWT_SECRET="short")
    assert json.loads("[]") == []
