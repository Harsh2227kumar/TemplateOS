"""
V1.7 Phase 1: GET /api/v1/documents/{document_id}/preview.

Coverage:
- 401 without auth, 403 for another user, 404 for a missing document.
- Never generated: generation is null, all permissions true for a draft.
- Shared formatting: a date's display_value is the exact text the generated
  DOCX shows (generation_service.format_field_value is used by both).
- List fields: items parsed, blanks dropped; an empty required list is in
  missing_required.
- Generated + PDF converted (mocked Gotenberg): availability flags and URLs.
- Generated file deleted from disk: docx_available false, still 200.
- is_stale flips after a later save.
- Status "submitted": can_edit / can_generate false, can_download true.
- The raw JSON never contains a storage path or an absolute path.
- The preview never generates or converts anything.

Storage is redirected to a temporary folder; the database is in-memory
SQLite with foreign keys enforced. PDF conversion is mocked (never a real
Gotenberg or soffice).
"""

import io
import json
import os
import re

import httpx
import pytest
from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

os.environ.setdefault("DATABASE_URL", "sqlite://")
os.environ.setdefault(
    "JWT_SECRET_KEY", "test-secret-that-is-long-enough-for-auth-tests"
)

from app.core.config import settings
from app.crud.template_field_crud import create_field
from app.db.base import Base
from app.db.session import get_db
from app.main import app
from app.models.document import Document
from app.models.document_value import DocumentValue
from app.models.generated_file import GeneratedFile
from app.models.template import Template
from app.models.template_field import TemplateField
from app.models.user import User
from app.schemas.template_field import TemplateFieldCreate
from app.services import pdf_service
from app.services.generation_service import build_context, format_field_value
from app.services.storage_service import storage_service

test_engine = create_engine(
    "sqlite://",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)


@event.listens_for(test_engine, "connect")
def _enable_foreign_keys(dbapi_connection, _record):
    cursor = dbapi_connection.cursor()
    cursor.execute("PRAGMA foreign_keys=ON")
    cursor.close()


TestSession = sessionmaker(bind=test_engine, autoflush=False, autocommit=False)


def override_get_db():
    with TestSession() as db:
        yield db


client = TestClient(app)

OWNER_EMAIL = "preview_owner@example.com"
STRANGER_EMAIL = "preview_stranger@example.com"
GOTENBERG_URL = "http://gotenberg.test:3000"
FAKE_PDF = b"%PDF-1.7\n% fake pdf for preview tests\n%%EOF"


@pytest.fixture(scope="module", autouse=True)
def setup_db(tmp_path_factory):
    old = app.dependency_overrides.get(get_db)
    old_root = storage_service.root_path
    storage_service.root_path = tmp_path_factory.mktemp("preview-storage").resolve()
    app.dependency_overrides[get_db] = override_get_db
    Base.metadata.drop_all(bind=test_engine)
    Base.metadata.create_all(bind=test_engine)
    yield
    Base.metadata.drop_all(bind=test_engine)
    storage_service.root_path = old_root
    if old is not None:
        app.dependency_overrides[get_db] = old
    else:
        app.dependency_overrides.pop(get_db, None)


@pytest.fixture(autouse=True)
def pdf_off(monkeypatch):
    """No PDF converter unless a test configures one (never a real soffice)."""
    monkeypatch.setattr(settings, "pdf_converter", "none")
    monkeypatch.setattr(settings, "gotenberg_url", None)
    monkeypatch.setattr(settings, "libreoffice_path", None)


class FakeGotenberg:
    def __init__(self) -> None:
        self.calls: list[httpx.Request] = []

    def __call__(self, request: httpx.Request) -> httpx.Response:
        request.read()
        self.calls.append(request)
        return httpx.Response(200, content=FAKE_PDF)


@pytest.fixture()
def gotenberg(monkeypatch) -> FakeGotenberg:
    fake = FakeGotenberg()
    monkeypatch.setattr(settings, "pdf_conversion_enabled", True)
    monkeypatch.setattr(settings, "pdf_converter", "auto")
    monkeypatch.setattr(settings, "gotenberg_url", GOTENBERG_URL)
    monkeypatch.setattr(settings, "gotenberg_timeout_seconds", 30)
    monkeypatch.setattr(settings, "libreoffice_path", None)
    monkeypatch.setattr(pdf_service, "gotenberg_transport", httpx.MockTransport(fake))
    return fake


def _signup_and_login(email: str, full_name: str) -> str:
    client.post(
        "/api/v1/auth/signup",
        json={"full_name": full_name, "email": email, "password": "strong-password"},
    )
    login = client.post(
        "/api/v1/auth/login", json={"email": email, "password": "strong-password"}
    )
    assert login.status_code == 200, login.text
    return login.json()["access_token"]


@pytest.fixture(scope="module")
def owner_token(setup_db):
    return _signup_and_login(OWNER_EMAIL, "Preview Owner")


@pytest.fixture(scope="module")
def stranger_token(setup_db):
    return _signup_and_login(STRANGER_EMAIL, "Preview Stranger")


def _auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


# ── Fixtures: template, draft, values ────────────────────────────

TEMPLATE_PARAGRAPHS = [
    "Title: {{ title }}",
    "Date: {{ meeting_date }}",
    "Count: {{ count }}",
    "Notes: {{ notes }}",
    "Agenda: {{ agenda_items }}",
    "Actions: {{ action_items }}",
]

# (field_name, field_type, is_required, section, display_order)
FIELDS = [
    ("title", "text", True, "Meeting Details", 1),
    ("meeting_date", "date", True, "Meeting Details", 2),
    ("count", "number", False, None, 3),
    ("notes", "textarea", False, "Minutes", 4),
    ("agenda_items", "list", True, "Minutes", 5),
    ("action_items", "list", True, "Minutes", 6),
]

GOOD_VALUES = {
    "title": "Budget review",
    "meeting_date": "2026-10-14",
    "count": "42",
    "agenda_items": json.dumps(["Welcome", " ", "Budget", ""]),
    "action_items": json.dumps(["Send minutes"]),
}


def _docx_bytes(paragraphs: list[str]) -> bytes:
    document = DocxDocument()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _owner_id() -> int:
    with TestSession() as session:
        return session.execute(
            select(User.id).where(User.email == OWNER_EMAIL)
        ).scalar_one()


def _make_template() -> int:
    stored = storage_service.save_bytes(
        "templates_processed", "preview.docx", _docx_bytes(TEMPLATE_PARAGRAPHS)
    )
    with TestSession() as session:
        template = Template(
            name="Preview Template",
            category="mom",
            visibility="public",
            uploaded_by=_owner_id(),
            status="field_configured",
            processed_file_path=stored.path,
        )
        session.add(template)
        session.commit()
        template_id = template.id
        for name, field_type, required, section, order in FIELDS:
            create_field(
                session,
                template_id,
                TemplateFieldCreate(
                    field_name=name,
                    field_label=name.replace("_", " ").title(),
                    field_type=field_type,
                    is_required=required,
                    section=section,
                    display_order=order,
                ),
            )
    return template_id


def _create_doc(token: str, template_id: int) -> int:
    response = client.post(
        "/api/v1/documents/create",
        headers=_auth(token),
        json={"template_id": template_id, "name": "Academic Committee MoM"},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _save(token: str, document_id: int, values: dict[str, str]) -> None:
    response = client.put(
        f"/api/v1/documents/{document_id}/values",
        headers=_auth(token),
        json={"values": [{"field_name": k, "value": v} for k, v in values.items()]},
    )
    assert response.status_code == 200, response.text


def _generate(token: str, document_id: int) -> dict:
    response = client.post(
        f"/api/v1/documents/{document_id}/generate", headers=_auth(token)
    )
    assert response.status_code == 200, response.text
    return response.json()


def _preview(token: str, document_id: int):
    return client.get(f"/api/v1/documents/{document_id}/preview", headers=_auth(token))


def _fields_by_name(body: dict) -> dict[str, dict]:
    return {f["field_name"]: f for s in body["sections"] for f in s["fields"]}


def _record(document_id: int) -> GeneratedFile | None:
    with TestSession() as session:
        return session.execute(
            select(GeneratedFile).where(GeneratedFile.document_id == document_id)
        ).scalar_one_or_none()


@pytest.fixture()
def draft(owner_token) -> int:
    """A draft with every value saved, never generated."""
    document_id = _create_doc(owner_token, _make_template())
    _save(owner_token, document_id, GOOD_VALUES)
    return document_id


# ── 1. Shared formatter ───────────────────────────────────────────


@pytest.mark.parametrize(
    "field_type, raw, expected",
    [
        ("date", "2026-10-14", "14-10-2026"),
        ("date", "2026-02-30", "2026-02-30"),
        ("date", "14/10/2026", "14/10/2026"),
        ("number", "42", "42"),
        ("text", "A & B", "A & B"),
        ("textarea", "line 1\r\nline 2", "line 1\nline 2"),
        ("list", json.dumps(["a", " ", "b"]), "a\nb"),
        ("list", "not json", "not json"),
        ("text", None, ""),
        ("text", "   ", ""),
        ("list", "[]", ""),
    ],
)
def test_format_field_value_matches_generation_context(field_type, raw, expected):
    field = TemplateField(field_name="x", field_type=field_type)
    assert format_field_value(field, raw) == expected
    if raw is not None and raw.strip():
        assert format_field_value(field, raw) == str(build_context([field], {"x": raw})["x"])


# ── 2. Errors ─────────────────────────────────────────────────────


def test_preview_requires_auth(draft):
    response = client.get(f"/api/v1/documents/{draft}/preview")
    assert response.status_code == 401


def test_preview_other_user_forbidden(draft, stranger_token):
    response = _preview(stranger_token, draft)
    assert response.status_code == 403
    assert response.json()["detail"] == "You can only preview your own documents"


def test_preview_missing_document(owner_token):
    response = _preview(owner_token, 999_999)
    assert response.status_code == 404
    assert response.json()["detail"] == "Document not found"


# ── 3. Draft, never generated ─────────────────────────────────────


def test_preview_draft_never_generated(owner_token, draft):
    response = _preview(owner_token, draft)
    assert response.status_code == 200, response.text
    body = response.json()

    assert body["document"]["id"] == draft
    assert body["document"]["name"] == "Academic Committee MoM"
    assert body["document"]["status"] == "draft"
    assert body["template"]["name"] == "Preview Template"
    assert body["template"]["category"] == "mom"
    assert body["template"]["version"] == 1
    assert body["template"]["has_processed_file"] is True
    assert body["generation"] is None
    assert body["permissions"] == {
        "can_edit": True,
        "can_generate": True,
        "can_download": True,
    }
    assert body["completeness"] == {
        "required_total": 4,
        "required_filled": 4,
        "missing_required": [],
    }
    assert [s["name"] for s in body["sections"]] == ["Meeting Details", "General", "Minutes"]

    fields = _fields_by_name(body)
    assert fields["meeting_date"]["value"] == "2026-10-14"
    assert fields["meeting_date"]["display_value"] == "14-10-2026"
    assert fields["meeting_date"]["items"] is None
    assert fields["count"]["display_value"] == "42"
    assert fields["notes"]["value"] is None
    assert fields["notes"]["display_value"] == ""
    assert fields["notes"]["is_empty"] is True
    assert fields["notes"]["missing_required"] is False
    assert fields["agenda_items"]["items"] == ["Welcome", "Budget"]
    assert fields["agenda_items"]["display_value"] == "Welcome\nBudget"
    assert fields["agenda_items"]["is_empty"] is False


def test_preview_never_generates_or_converts(owner_token, draft, gotenberg):
    assert _preview(owner_token, draft).status_code == 200
    assert _record(draft) is None
    assert gotenberg.calls == []


def test_empty_required_list_is_missing(owner_token):
    document_id = _create_doc(owner_token, _make_template())
    values = {k: v for k, v in GOOD_VALUES.items() if k != "action_items"}
    _save(owner_token, document_id, values)
    # The save validator rejects a blank required list, so write it directly
    # (as an older save or an import could have).
    with TestSession() as session:
        session.add(
            DocumentValue(
                document_id=document_id,
                field_name="action_items",
                value=json.dumps(["", " "]),
            )
        )
        session.commit()

    body = _preview(owner_token, document_id).json()
    action_items = _fields_by_name(body)["action_items"]
    assert action_items["items"] == []
    assert action_items["display_value"] == ""
    assert action_items["is_empty"] is True
    assert action_items["missing_required"] is True
    assert body["completeness"]["required_total"] == 4
    assert body["completeness"]["required_filled"] == 3
    assert body["completeness"]["missing_required"] == ["action_items"]


# ── 4. Generated ──────────────────────────────────────────────────


def test_display_value_matches_generated_docx(owner_token, draft):
    _generate(owner_token, draft)
    body = _preview(owner_token, draft).json()
    fields = _fields_by_name(body)

    docx = DocxDocument(io.BytesIO(storage_service.read_bytes(_record(draft).docx_path)))
    paragraphs = [p.text for p in docx.paragraphs]
    assert fields["meeting_date"]["display_value"] == "14-10-2026"
    assert f"Date: {fields['meeting_date']['display_value']}" in paragraphs
    assert f"Title: {fields['title']['display_value']}" in paragraphs
    assert f"Count: {fields['count']['display_value']}" in paragraphs
    assert f"Agenda: {fields['agenda_items']['display_value']}" in paragraphs


def test_generated_with_pdf_converted(owner_token, draft, gotenberg):
    generated = _generate(owner_token, draft)
    assert generated["pdf_status"] == "converted"
    calls_after_generate = len(gotenberg.calls)

    response = _preview(owner_token, draft)
    assert response.status_code == 200, response.text
    generation = response.json()["generation"]
    assert generation["generated_at"] == generated["generated_at"]
    assert generation["is_stale"] is False
    assert generation["template_version"] == 1
    assert generation["docx_available"] is True
    assert generation["docx_url"] == f"/api/v1/documents/{draft}/download?format=docx"
    assert generation["pdf_status"] == "converted"
    assert generation["pdf_error"] is None
    assert generation["pdf_available"] is True
    assert generation["pdf_url"] == f"/api/v1/documents/{draft}/download/pdf"
    assert generation["pdf_filename"] == generated["pdf_filename"]
    assert generation["pdf_filename"].endswith(".pdf")
    assert "/" not in generation["pdf_filename"]
    assert len(gotenberg.calls) == calls_after_generate
    assert response.json()["permissions"]["can_edit"] is True  # status "generated"


def test_pdf_failed_has_no_pdf_url(owner_token, draft):
    _generate(owner_token, draft)  # converter disabled -> pdf_status "failed"
    generation = _preview(owner_token, draft).json()["generation"]
    assert generation["pdf_status"] == "failed"
    assert generation["pdf_error"]
    assert generation["pdf_available"] is False
    assert generation["pdf_url"] is None
    assert generation["pdf_filename"] is None
    assert generation["docx_available"] is True


def test_files_deleted_from_disk_are_unavailable(owner_token, draft, gotenberg, caplog):
    _generate(owner_token, draft)
    record = _record(draft)
    storage_service.delete_file(record.docx_path)
    storage_service.delete_file(record.pdf_path)

    with caplog.at_level("WARNING"):
        response = _preview(owner_token, draft)
    assert response.status_code == 200, response.text
    generation = response.json()["generation"]
    assert generation["docx_available"] is False
    assert generation["pdf_available"] is False
    assert generation["pdf_url"] is None
    assert generation["pdf_status"] == "converted"  # the DB row is unchanged
    assert f"document {draft}" in caplog.text
    assert str(storage_service.root_path) not in caplog.text


def test_path_outside_storage_is_unavailable(owner_token, draft):
    _generate(owner_token, draft)
    with TestSession() as session:
        record = session.execute(
            select(GeneratedFile).where(GeneratedFile.document_id == draft)
        ).scalar_one()
        record.docx_path = "../../outside.docx"
        session.commit()
    response = _preview(owner_token, draft)
    assert response.status_code == 200, response.text
    assert response.json()["generation"]["docx_available"] is False


def test_is_stale_after_later_save(owner_token, draft):
    _generate(owner_token, draft)
    assert _preview(owner_token, draft).json()["generation"]["is_stale"] is False

    _save(owner_token, draft, {"title": "Budget review (amended)"})
    body = _preview(owner_token, draft).json()
    assert body["generation"]["is_stale"] is True
    assert _fields_by_name(body)["title"]["display_value"] == "Budget review (amended)"


def test_submitted_document_is_read_only(owner_token, draft):
    _generate(owner_token, draft)
    with TestSession() as session:
        session.get(Document, draft).status = "submitted"
        session.commit()

    body = _preview(owner_token, draft).json()
    assert body["document"]["status"] == "submitted"
    assert body["permissions"] == {
        "can_edit": False,
        "can_generate": False,
        "can_download": True,
    }


def test_response_never_exposes_storage_paths(owner_token, draft, gotenberg):
    _generate(owner_token, draft)
    record = _record(draft)
    raw = _preview(owner_token, draft).text

    root = str(storage_service.root_path)
    assert root not in raw
    assert root.replace("\\", "/") not in raw
    assert json.dumps(root)[1:-1] not in raw  # JSON-escaped backslashes
    assert "storage/" not in raw
    assert record.docx_path not in raw
    assert record.pdf_path not in raw
    assert "generated/docx" not in raw and "generated/pdf" not in raw
    assert "templates/processed" not in raw
    assert not re.search(r"[A-Za-z]:(\\\\|/)", raw)  # no Windows drive paths
    for value in json.loads(raw)["generation"].values():
        if isinstance(value, str):
            assert not value.startswith("/") or value.startswith(settings.api_v1_prefix)
